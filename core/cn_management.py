from __future__ import annotations

from datetime import date, timedelta
from io import BytesIO
import logging
from typing import Any
import zipfile

import openpyxl
import pandas as pd
import plotly.express as px
import plotly.graph_objects as go
import streamlit as st
from sqlalchemy.exc import SQLAlchemyError

from database.calibration_store import (
    PAPEIS_AUDITORIA,
    _engine,
    autenticar_usuario,
    criar_usuario,
    exportar_dataframe_csv,
    exportar_dataframe_excel,
    listar_usuarios,
    usuario_ativo,
)
from database.cn_store import (
    CRITICIDADES,
    ETAPAS_CN,
    FAMILIA_REFERENCIA_GLOBAL,
    atualizar_cn,
    atualizar_etapa,
    baixar_backup_cns,
    carregar_auditoria,
    carregar_auditoria_administracao,
    carregar_cns,
    carregar_notificacoes,
    carregar_referencias_cns,
    carregar_responsaveis,
    carregar_todas_notificacoes_administracao,
    confirmar_recebimento_etapa,
    criar_cn,
    criar_backup_cns,
    data_para_iso,
    excluir_cn,
    importar_cns,
    administrar_responsavel,
    inicializar_banco,
    listar_numeros_cns_importadas,
    listar_responsaveis_administracao,
    limpar_cns_importadas,
    limpar_dados_cns,
    listar_backups_cns,
    marcar_notificacao_lida,
    processar_notificacoes_cns,
    registrar_alerta_importacao,
    reimportar_base_cns,
    salvar_referencias_cns,
    salvar_responsavel,
)


logger = logging.getLogger(__name__)

_COLUNAS_ETAPAS = dict(zip(ETAPAS_CN, "GHIJKLMNOP"))
_COLUNAS_IMPORTACAO = {
    "numero_cn": 0,
    "familia": 1,
    "descricao": 2,
    "codigo_linha": 3,
    "codigo_cn": 4,
    "resumo_modificacao": 5,
    "observacao": 16,
    "status_legado": 17,
}


def _valor_texto(valor: Any) -> str:
    return "" if valor is None else str(valor).strip()


def analisar_planilha_cns(conteudo: bytes) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    try:
        workbook = openpyxl.load_workbook(
            BytesIO(conteudo),
            data_only=True,
            read_only=True,
        )
    except (
        OSError,
        ValueError,
        zipfile.BadZipFile,
        openpyxl.utils.exceptions.InvalidFileException,
    ) as erro:
        raise ValueError("O arquivo não é uma planilha Excel válida.") from erro
    try:
        nome_controle = "CONTROLE DE CN'S 2026"
        if nome_controle not in workbook.sheetnames:
            raise ValueError(f"A planilha deve conter a aba '{nome_controle}'.")
        if "Planilha1" not in workbook.sheetnames:
            raise ValueError("A planilha deve conter a aba de referência 'Planilha1'.")
        ws = workbook[nome_controle]
        apoio = workbook["Planilha1"]
        linhas_apoio = list(
            apoio.iter_rows(min_row=1, max_col=4, values_only=True)
        )
        referencias_planilha = [
            {
                "familia": _valor_texto(linha[0]),
                "descricao": _valor_texto(linha[2]),
                "codigo_linha": _valor_texto(linha[3]),
            }
            for linha in linhas_apoio
            if any(_valor_texto(linha[indice]) for indice in (0, 2, 3))
        ]
        st.session_state["_referencias_cns"] = {
            "familias": sorted(
                {
                    _valor_texto(linha[0])
                    for linha in linhas_apoio
                    if _valor_texto(linha[0])
                }
            ),
            "descricoes": sorted(
                {
                    _valor_texto(linha[2])
                    for linha in linhas_apoio
                    if _valor_texto(linha[2])
                }
            ),
            "codigo_linha": {
                _valor_texto(linha[2]).casefold(): _valor_texto(linha[3])
                for linha in linhas_apoio
                if _valor_texto(linha[2])
            },
            "linhas": referencias_planilha,
        }
        linha_por_descricao = {
            _valor_texto(linha[2]).casefold(): _valor_texto(linha[3])
            for linha in linhas_apoio
            if _valor_texto(linha[2])
        }
        linha_por_familia_descricao = {
            (
                _valor_texto(linha[0]).casefold(),
                _valor_texto(linha[2]).casefold(),
            ): _valor_texto(linha[3])
            for linha in linhas_apoio
            if _valor_texto(linha[0]) and _valor_texto(linha[2])
        }
        registros: list[dict[str, Any]] = []
        alertas: list[dict[str, Any]] = []
        for numero_linha, linha in enumerate(
            ws.iter_rows(min_row=3, max_col=18, values_only=True),
            start=3,
        ):
            valores = list(linha)
            if not any(valor is not None for valor in valores):
                continue
            codigo = _valor_texto(valores[_COLUNAS_IMPORTACAO["codigo_cn"]])
            if not codigo:
                continue
            descricao = _valor_texto(valores[_COLUNAS_IMPORTACAO["descricao"]])
            registro: dict[str, Any] = {
                campo: _valor_texto(valores[indice])
                for campo, indice in _COLUNAS_IMPORTACAO.items()
                if campo not in {"status_legado"}
            }
            registro["numero_cn"] = _valor_texto(
                valores[_COLUNAS_IMPORTACAO["numero_cn"]]
            ) or codigo
            registro["codigo_linha"] = linha_por_familia_descricao.get(
                (
                    registro["familia"].casefold(),
                    descricao.casefold(),
                ),
                linha_por_descricao.get(descricao.casefold(), ""),
            )
            if not registro["codigo_linha"]:
                alertas.append(
                    {
                        "numero_cn": registro["numero_cn"],
                        "coluna": "D",
                        "linha": numero_linha,
                        "detalhe": (
                            f"Linha {numero_linha}, coluna D: descrição "
                            f"'{descricao}' não encontrada na Planilha1 coluna C."
                        ),
                    }
                )
            registro["responsavel_email"] = ""
            datas: dict[str, str] = {}
            for etapa, letra in _COLUNAS_ETAPAS.items():
                indice_coluna = ord(letra) - ord("A")
                data, alerta = data_para_iso(valores[indice_coluna], letra, numero_linha)
                datas[etapa] = data
                if alerta:
                    alertas.append(
                        {
                            "numero_cn": registro["numero_cn"],
                            "coluna": letra,
                            "linha": numero_linha,
                            "detalhe": alerta,
                        }
                    )
            registro["etapas"] = datas
            registros.append(registro)
        return registros, alertas
    finally:
        workbook.close()


def _usuario_autenticado() -> dict[str, Any] | None:
    sessao = st.session_state.get("_usuario_gestao_cns")
    if not isinstance(sessao, dict):
        return None
    try:
        usuario = usuario_ativo(sessao["email"])
    except SQLAlchemyError as erro:
        logger.exception("Falha ao validar sessão da Gestão de CN's.")
        st.error(f"Não foi possível validar sua sessão: {erro}")
        return None
    if usuario is None:
        st.session_state.pop("_usuario_gestao_cns", None)
        return None
    st.session_state["_usuario_gestao_cns"] = usuario
    return usuario


def _render_autenticacao(titulo: str = "Gestão de CN's") -> dict[str, Any] | None:
    usuario = _usuario_autenticado()
    if usuario is not None:
        return usuario
    st.subheader(f"Acesso à {titulo}")
    entrar, solicitar = st.tabs(["Entrar", "Solicitar acesso"])
    with entrar:
        with st.form("login_gestao_cns"):
            email = st.text_input("E-mail", key="login_cns_email")
            senha = st.text_input("Senha", type="password", key="login_cns_senha")
            enviado = st.form_submit_button("Entrar", type="primary")
        if enviado:
            try:
                usuario = autenticar_usuario(email, senha)
            except (ValueError, SQLAlchemyError) as erro:
                logger.exception("Falha na autenticação da Gestão de CN's.")
                st.error(f"Não foi possível autenticar: {erro}")
            else:
                if usuario is None:
                    st.error("Credenciais inválidas ou solicitação ainda não aprovada.")
                else:
                    st.session_state["_usuario_gestao_cns"] = usuario
                    st.rerun()
    with solicitar:
        with st.form("solicitar_acesso_cns"):
            nome = st.text_input("Nome completo", key="solicitar_cns_nome")
            novo_email = st.text_input("E-mail corporativo", key="solicitar_cns_email")
            senha_nova = st.text_input(
                "Senha (mínimo 12 caracteres)",
                type="password",
                key="solicitar_cns_senha",
            )
            confirmar = st.text_input(
                "Confirme a senha",
                type="password",
                key="solicitar_cns_confirmar_senha",
            )
            enviado = st.form_submit_button("Enviar solicitação")
        if enviado:
            if senha_nova != confirmar:
                st.error("As senhas não coincidem.")
            else:
                try:
                    criar_usuario(nome, novo_email, senha_nova)
                except (ValueError, SQLAlchemyError) as erro:
                    logger.exception("Falha ao solicitar acesso à Gestão de CN's.")
                    st.error(f"Não foi possível solicitar acesso: {erro}")
                else:
                    st.success("Solicitação registrada para aprovação.")
    try:
        usuarios = listar_usuarios()
    except SQLAlchemyError as erro:
        logger.exception("Falha ao consultar usuários da Gestão de CN's.")
        st.error(f"Não foi possível consultar usuários: {erro}")
    else:
        if usuarios.empty:
            st.info(
                "Ainda não existe administrador. No ambiente local, crie o primeiro "
                "administrador com `python -m database.provision_calibration_admin`."
            )
    return None


def _formatar_data(valor: Any) -> date | None:
    if valor is None or pd.isna(valor) or str(valor).upper() == "NA":
        return None
    try:
        return date.fromisoformat(str(valor)[:10])
    except ValueError:
        return None


def _formatar_identificador_cn(valor: Any) -> str:
    texto = _valor_texto(valor)
    if not texto:
        return "CN #—"
    if texto.isdigit():
        return f"CN #{int(texto):03d}"
    if texto.upper().startswith("CN #"):
        return texto
    return f"CN #{texto}"


def _normalizar_status_visual(valor: Any) -> str:
    texto = _valor_texto(valor).upper()
    return {
        "IMPLEMENTADA": "IMPLEMENTADO",
        "PROGRAMADA": "PROGRAMADO",
        "ATRASADA": "ATRASADO",
    }.get(texto, texto)


def _contagens_status(dataframe: pd.DataFrame) -> dict[str, int]:
    status = dataframe["status"].map(_normalizar_status_visual)
    return {
        "IMPLEMENTADO": int(status.eq("IMPLEMENTADO").sum()),
        "PROGRAMADO": int(status.eq("PROGRAMADO").sum()),
        "ATRASADO": int(status.eq("ATRASADO").sum()),
        "OUTROS": int(
            (~status.isin(("IMPLEMENTADO", "PROGRAMADO", "ATRASADO"))).sum()
        ),
    }


def _figura_contagem(dataframe: pd.DataFrame, coluna: str, titulo: str) -> Any:
    contagem = dataframe[coluna].fillna("Não informado").astype(str).value_counts()
    contagem = contagem.head(8).sort_values()
    figura = px.bar(
        x=contagem.values,
        y=contagem.index,
        orientation="h",
        title=titulo,
        text=contagem.values,
    )
    figura.update_traces(
        marker_color="#4da3ff",
        textposition="outside",
        hovertemplate="%{y}: %{x} CNs<extra></extra>",
    )
    figura.update_layout(
        paper_bgcolor="rgba(0,0,0,0)",
        plot_bgcolor="rgba(0,0,0,0)",
        font_color="#f7f9fc",
        showlegend=False,
        xaxis_title="Quantidade de CNs",
        yaxis_title=None,
        height=max(250, len(contagem) * 34 + 90),
        margin=dict(l=15, r=35, t=50, b=20),
    )
    return figura


def _indicadores_dashboard(dataframe: pd.DataFrame) -> dict[str, int]:
    contagens = _contagens_status(dataframe)
    pendentes = dataframe["status_confirmacao"].fillna("PENDENTE").astype(str).str.strip().str.upper().eq("PENDENTE")
    criticidade = dataframe["criticidade"].fillna("").astype(str).str.strip().str.casefold()
    status = dataframe["status"].map(_normalizar_status_visual)
    alta = criticidade.eq("alta")
    return {
        "total": len(dataframe),
        "implementadas": contagens["IMPLEMENTADO"],
        "programadas": contagens["PROGRAMADO"],
        "atrasadas": contagens["ATRASADO"],
        "pendentes_confirmacao": int(pendentes.sum()),
        "criticas": int(alta.sum()),
        "criticas_atrasadas": int((alta & status.eq("ATRASADO")).sum()),
        "sem_status": contagens["OUTROS"],
    }


def _gerar_dashboard(dataframe: pd.DataFrame) -> None:
    st.subheader("Dashboard executivo")
    if dataframe.empty:
        st.info("Cadastre uma CN ou importe a planilha para iniciar o acompanhamento.")
        return
    indicadores = _indicadores_dashboard(dataframe)
    kpis = (
        ("Total CNs", indicadores["total"]),
        ("🟢 Implementadas", indicadores["implementadas"]),
        ("🟡 Programadas", indicadores["programadas"]),
        ("🔴 Atrasadas", indicadores["atrasadas"]),
        ("Pendentes de confirmação", indicadores["pendentes_confirmacao"]),
        ("CNs críticas", indicadores["criticas"]),
    )
    for inicio in (0, 3):
        colunas = st.columns(3)
        for coluna, (titulo, valor) in zip(colunas, kpis[inicio : inicio + 3]):
            with coluna.container(border=True):
                st.metric(titulo, valor)
                if titulo == "CNs críticas":
                    st.caption(
                        f"{indicadores['criticas_atrasadas']} crítica(s) atrasada(s)"
                    )
    soma_status = (
        indicadores["implementadas"]
        + indicadores["programadas"]
        + indicadores["atrasadas"]
    )
    st.caption(
        f"Status classificados: {soma_status} de {indicadores['total']} CNs"
        + (
            f" · {indicadores['sem_status']} sem status reconhecido"
            if indicadores["sem_status"]
            else ""
        )
    )

    registros = dataframe.to_dict("records")
    etapas_atrasadas = [
        {
            "CN": _formatar_identificador_cn(registro["numero_cn"]),
            "Etapa": etapa["etapa"],
        }
        for registro in registros
        for etapa in (registro.get("etapas") or [])
        if etapa.get("status_calculado") == "ATRASADO"
    ]
    resumo_colunas = st.columns(2)
    with resumo_colunas[0]:
        st.plotly_chart(
            _figura_contagem(dataframe, "familia", "CNs por família"),
            use_container_width=True,
        )
    with resumo_colunas[1]:
        st.subheader("Etapas com mais atrasos")
        atrasos = pd.DataFrame(etapas_atrasadas)
        if atrasos.empty:
            st.success("Nenhuma etapa atrasada.")
        else:
            contagem_etapas = atrasos["Etapa"].value_counts().rename_axis(
                "Etapa"
            ).reset_index(name="Atrasos")
            st.dataframe(
                contagem_etapas.head(5),
                hide_index=True,
                use_container_width=True,
                height=min(300, 44 + 36 * min(5, len(contagem_etapas))),
            )

    lead_times: list[int] = []
    aderencias: list[bool] = []
    implementacoes: list[dict[str, str]] = []
    for registro in registros:
        etapas_registro = registro.get("etapas") or []
        concluidas = [
            etapa
            for etapa in etapas_registro
            if etapa.get("data_conclusao")
        ]
        if concluidas:
            data_criacao = pd.to_datetime(registro["data_criacao"], errors="coerce")
            data_final = max(
                pd.to_datetime(etapa["data_conclusao"], errors="coerce")
                for etapa in concluidas
            )
            if pd.notna(data_criacao) and pd.notna(data_final):
                lead_times.append((data_final - data_criacao).days)
        if _normalizar_status_visual(registro["status"]) == "IMPLEMENTADO" and concluidas:
            implementacoes.append(
                {
                    "Mês": max(
                        pd.to_datetime(etapa["data_conclusao"], errors="coerce")
                        for etapa in concluidas
                    ).to_period("M").strftime("%Y-%m")
                }
            )
        for etapa in concluidas:
            planejada = _formatar_data(etapa.get("data_planejada"))
            realizada = _formatar_data(etapa.get("data_conclusao"))
            if planejada is not None and realizada is not None:
                aderencias.append(realizada <= planejada)
    with st.expander("Indicadores de prazo e tendência"):
        metricas = st.columns(3)
        media_lead = sum(lead_times) / len(lead_times) if lead_times else None
        media_aderencia = (
            sum(aderencias) / len(aderencias) * 100 if aderencias else None
        )
        metricas[0].metric(
            "Lead Time Médio",
            f"{media_lead:.1f} dias" if media_lead is not None else "—",
        )
        metricas[1].metric(
            "Aderência ao prazo",
            f"{media_aderencia:.1f}%" if media_aderencia is not None else "—",
        )
        metricas[2].metric(
            "CNs críticas atrasadas", indicadores["criticas_atrasadas"]
        )
        if implementacoes:
            st.plotly_chart(
                _figura_contagem(
                    pd.DataFrame(implementacoes),
                    "Mês",
                    "CNs implementadas por mês",
                ),
                use_container_width=True,
            )


def _dados_gantt_cns(dataframe: pd.DataFrame, ano: int) -> pd.DataFrame:
    tarefas: list[dict[str, Any]] = []
    for registro in dataframe.to_dict("records"):
        etapas_por_nome = {
            etapa["etapa"]: etapa for etapa in registro.get("etapas") or []
        }
        criacao = _formatar_data(registro.get("data_criacao"))
        data_anterior = criacao
        for nome_etapa in ETAPAS_CN:
            etapa = etapas_por_nome.get(nome_etapa)
            if etapa is None:
                continue
            prazo = _formatar_data(etapa.get("data_planejada"))
            conclusao = _formatar_data(
                etapa.get("data_conclusao") or etapa.get("data_recebimento")
            )
            if prazo is not None:
                if prazo.year == ano:
                    inicio = data_anterior or date(prazo.year, 1, 1)
                    if inicio >= prazo:
                        inicio = prazo - timedelta(days=1)
                    tarefas.append(
                        {
                            "CN": _formatar_identificador_cn(registro["numero_cn"]),
                            "Etapa": nome_etapa,
                            "Atividade": (
                                f"{_formatar_identificador_cn(registro['numero_cn'])}"
                                f" · {nome_etapa}"
                            ),
                            "Início": inicio,
                            "Prazo": prazo,
                            "Conclusão real": conclusao,
                            "Família": registro.get("familia") or "—",
                            "Responsável": etapa.get("responsavel_email") or "—",
                            "Status": etapa.get("status_calculado") or "PENDENTE",
                        }
                    )
                data_anterior = prazo
    colunas = (
        "CN",
        "Etapa",
        "Atividade",
        "Início",
        "Prazo",
        "Conclusão real",
        "Família",
        "Responsável",
        "Status",
    )
    return pd.DataFrame(tarefas, columns=colunas).sort_values(
        ["Prazo", "CN", "Etapa"],
        ignore_index=True,
    )


def _render_gantt_anual(dataframe: pd.DataFrame) -> None:
    st.subheader("Cronograma anual das CNs")
    if dataframe.empty:
        st.info("Cadastre uma CN com prazos planejados para montar o Gantt.")
        return

    anos_disponiveis: set[int] = set()
    for registro in dataframe.to_dict("records"):
        for etapa in registro.get("etapas") or []:
            prazo = _formatar_data(etapa.get("data_planejada"))
            if prazo is not None:
                anos_disponiveis.add(prazo.year)
    anos_disponiveis.add(date.today().year)
    anos = sorted(anos_disponiveis)
    ano_atual = date.today().year
    ano = st.selectbox(
        "Ano do cronograma",
        anos,
        index=anos.index(ano_atual) if ano_atual in anos else len(anos) - 1,
        key="gantt_cns_ano",
    )
    tarefas = _dados_gantt_cns(dataframe, ano)
    if tarefas.empty:
        st.info(f"Não há etapas com prazo planejado em {ano}.")
        return

    col_filtro, col_indicador = st.columns([2, 3])
    pesquisa = col_filtro.text_input(
        "Filtrar CN",
        placeholder="Digite o número da CN",
        key="gantt_cns_pesquisa",
    ).strip().casefold()
    status_opcoes = ["Todos", *sorted(tarefas["Status"].dropna().unique())]
    status = col_filtro.selectbox(
        "Status das etapas",
        status_opcoes,
        key="gantt_cns_status",
    )
    if pesquisa:
        tarefas = tarefas[
            tarefas["CN"].str.casefold().str.contains(pesquisa, regex=False)
        ]
    if status != "Todos":
        tarefas = tarefas[tarefas["Status"] == status]
    col_indicador.metric(
        "Etapas no cronograma",
        len(tarefas),
        f"{tarefas['CN'].nunique()} CN(s)" if not tarefas.empty else "Nenhuma CN",
    )
    if tarefas.empty:
        st.info("Nenhuma etapa corresponde aos filtros selecionados.")
        return


    st.caption(
        "Cada barra mostra a janela entre a data de criação (ou o prazo da etapa "
        "anterior) e o prazo planejado desta etapa. O losango marca a conclusão real."
    )
    st.plotly_chart(_figura_gantt_cns(tarefas, ano), use_container_width=True)


def _figura_gantt_cns(tarefas: pd.DataFrame, ano: int) -> go.Figure:
    figura = px.timeline(
        tarefas,
        x_start="Início",
        x_end="Prazo",
        y="Atividade",
        color="Status",
        color_discrete_map={
            "CONCLUÍDO": "#36a269",
            "EM ANDAMENTO": "#e5a33d",
            "PENDENTE": "#4da3ff",
            "ATRASADO": "#e05252",
        },
        hover_data={
            "CN": True,
            "Etapa": True,
            "Família": True,
            "Responsável": True,
            "Início": "|%d/%m/%Y",
            "Prazo": "|%d/%m/%Y",
            "Conclusão real": "|%d/%m/%Y",
            "Status": True,
            "Atividade": False,
        },
        title=f"Etapas planejadas em {ano}",
    )
    concluidas = tarefas[tarefas["Conclusão real"].notna()]
    if not concluidas.empty:
        figura.add_trace(
            go.Scatter(
                x=concluidas["Conclusão real"],
                y=concluidas["Atividade"],
                mode="markers",
                name="Conclusão real",
                marker={
                    "symbol": "diamond",
                    "size": 10,
                    "color": "#f4d06f",
                    "line": {"color": "#172b45", "width": 1},
                },
                customdata=concluidas[["CN", "Etapa"]],
                hovertemplate=(
                    "%{customdata[0]} · %{customdata[1]}"
                    "<br>Concluída em %{x|%d/%m/%Y}<extra></extra>"
                ),
            )
        )
    figura.update_layout(
        paper_bgcolor="rgba(0,0,0,0)",
        plot_bgcolor="rgba(0,0,0,0)",
        font_color="#f7f9fc",
        height=max(420, min(1000, 30 * len(tarefas) + 140)),
        margin=dict(l=15, r=25, t=60, b=25),
        legend_title_text="Situação",
        xaxis_title="Data",
        yaxis_title=None,
    )
    figura.update_xaxes(
        range=[date(ano, 1, 1), date(ano + 1, 1, 1)],
        dtick="M1",
        tickformat="%b",
        showgrid=True,
        gridcolor="rgba(148,163,184,.2)",
    )
    figura.update_yaxes(autorange="reversed")
    return figura


def _filtrar_cns(dataframe: pd.DataFrame) -> pd.DataFrame:
    if dataframe.empty:
        return dataframe
    st.subheader("Localizar CN")
    pesquisa = st.text_input(
        "Pesquisa inteligente",
        placeholder="Número, descrição, família, responsável, resumo ou observação",
        key="pesquisa_inteligente_cns",
    ).strip().casefold()
    filtrado = dataframe
    if pesquisa:
        searchable = (
            "numero_cn",
            "codigo_cn",
            "familia",
            "descricao",
            "responsavel_email",
            "resumo_modificacao",
            "observacao",
            "criticidade",
        )
        mascara = pd.Series(False, index=filtrado.index)
        for campo in searchable:
            mascara |= filtrado[campo].fillna("").astype(str).str.casefold().str.contains(
                pesquisa, regex=False
            )
        mascara |= filtrado["status"].map(_normalizar_status_visual).str.casefold().str.contains(
            pesquisa, regex=False
        )
        identificadores = filtrado["numero_cn"].map(
            _formatar_identificador_cn
        ).str.casefold()
        mascara |= identificadores.str.contains(pesquisa, regex=False)
        filtrado = filtrado[mascara]

    controles = st.columns(3)
    numero = controles[0].text_input(
        "Número CN",
        placeholder="Ex.: 125",
        key="filtro_numero_cns",
    ).strip().casefold().replace("cn #", "").lstrip("0")
    if numero:
        filtrado = filtrado[
            filtrado["numero_cn"].map(
                lambda valor: _valor_texto(valor).casefold().replace("cn #", "").lstrip("0")
            ).str.contains(numero, regex=False)
        ]
    for controle, rotulo, campo, chave in (
        (controles[1], "Família", "familia", "filtro_cns_familia"),
        (controles[2], "Descrição", "descricao", "filtro_cns_descricao"),
    ):
        valores = sorted(dataframe[campo].dropna().astype(str).unique())
        escolhido = controle.selectbox(
            rotulo, ["Todas", *valores], key=chave
        )
        if escolhido != "Todas":
            filtrado = filtrado[filtrado[campo].astype(str) == escolhido]
    controles = st.columns(3)
    status = controles[0].selectbox(
        "Status",
        ["Todos", "PROGRAMADO", "IMPLEMENTADO", "ATRASADO"],
        key="filtro_cns_status",
    )
    if status != "Todos":
        filtrado = filtrado[
            filtrado["status"].map(_normalizar_status_visual) == status
        ]
    responsaveis = sorted(dataframe["responsavel_email"].dropna().astype(str).unique())
    responsavel = controles[1].selectbox(
        "Responsável",
        ["Todos", *responsaveis],
        key="filtro_cns_responsavel",
    )
    if responsavel != "Todos":
        filtrado = filtrado[filtrado["responsavel_email"] == responsavel]
    criticidades = sorted(dataframe["criticidade"].dropna().astype(str).unique())
    criticidade = controles[2].selectbox(
        "Criticidade",
        ["Todas", *criticidades],
        key="filtro_cns_criticidade",
    )
    if criticidade != "Todas":
        filtrado = filtrado[filtrado["criticidade"] == criticidade]
    return filtrado


def _dataframe_exibicao(dataframe: pd.DataFrame) -> pd.DataFrame:
    status_icon = {
        "IMPLEMENTADO": "🟢 Implementado",
        "PROGRAMADO": "🟡 Programado",
        "ATRASADO": "🔴 Atrasado",
    }
    responsaveis = st.session_state.get("_responsaveis_cns_nomes", {})
    return pd.DataFrame(
        {
            "Número CN": dataframe["numero_cn"].map(_formatar_identificador_cn),
            "Descrição": dataframe["descricao"],
            "Família": dataframe["familia"],
            "Status": dataframe["status"].map(
                lambda valor: status_icon.get(
                    _normalizar_status_visual(valor), _valor_texto(valor)
                )
            ),
            "Responsável": dataframe["responsavel_email"].map(
                lambda email: responsaveis.get(email, email)
            ),
            "Criticidade": dataframe["criticidade"],
        },
        index=dataframe.index,
    )


def _referencias_para_familia(
    linhas: list[dict[str, Any]],
    familia: str,
) -> list[dict[str, Any]]:
    return [
        linha
        for linha in linhas
        if linha["familia"] in {familia, FAMILIA_REFERENCIA_GLOBAL}
    ]


def _codigo_linha_para_familia(
    linhas: list[dict[str, Any]],
    familia: str,
    descricao: str,
) -> str:
    linha_especifica = next(
        (
            linha["codigo_linha"]
            for linha in linhas
            if linha["familia"] == familia and linha["descricao"] == descricao
        ),
        None,
    )
    if linha_especifica is not None:
        return linha_especifica
    return next(
        (
            linha["codigo_linha"]
            for linha in linhas
            if linha["familia"] == FAMILIA_REFERENCIA_GLOBAL
            and linha["descricao"] == descricao
        ),
        "",
    )


def _tela_cadastro(
    usuario: dict[str, Any],
    responsaveis: pd.DataFrame,
    referencias: dict[str, Any],
) -> None:
    st.subheader("Cadastrar CN")
    if responsaveis.empty:
        st.warning("Cadastre um responsável ativo antes de criar uma CN.")
        return
    opcoes_resp = {
        f"{linha.usuario} · {linha.email} · {linha.setor}": linha.email
        for linha in responsaveis.itertuples(index=False)
    }
    linhas_validas = [
        linha
        for linha in referencias["linhas"]
        if linha["familia"] and linha["descricao"]
    ]
    familias = sorted(
        {
            linha["familia"]
            for linha in linhas_validas
            if linha["familia"] != FAMILIA_REFERENCIA_GLOBAL
        }
    )
    if not familias:
        st.info(
            "Cadastre Família e Descrição em ⚙ Administração ou atualize esses "
            "cadastros pela importação da planilha antes de criar uma CN."
        )
        return
    familia_col, descricao_col = st.columns(2)
    familia = familia_col.selectbox(
        "Família *",
        familias,
        key="nova_cn_familia",
    )
    linhas_familia = _referencias_para_familia(linhas_validas, familia)
    descricoes = sorted({linha["descricao"] for linha in linhas_familia})
    descricao = descricao_col.selectbox(
        "Descrição *",
        descricoes,
        key="nova_cn_descricao",
    )
    codigo_linha = _codigo_linha_para_familia(
        linhas_familia,
        familia,
        descricao,
    )
    with st.form("form_cadastro_cn", clear_on_submit=True):
        c1, c2, c3 = st.columns(3)
        numero = c1.text_input("Número CN *")
        codigo = c2.text_input("Código CN *")
        c2.text_input(
            "Código Linha (automático)",
            value=codigo_linha,
            disabled=True,
            help="Derivado da descrição selecionada no cadastro mestre.",
        )
        responsavel_label = c3.selectbox("Responsável *", list(opcoes_resp))
        resumo = st.text_area("Resumo da Modificação")
        observacao = st.text_area("Observação")
        criticidade = st.selectbox("Criticidade", CRITICIDADES, index=1)
        st.caption("Informe somente prazos. A conclusão real deve ser registrada separadamente no acompanhamento.")
        etapas: dict[str, dict[str, Any]] = {}
        with st.expander("Prazos das etapas"):
            for etapa in ETAPAS_CN:
                d1, d2 = st.columns([1, 2])
                prazo = d1.date_input(
                    f"{etapa} — prazo (opcional)",
                    value=None,
                    key=f"novo_prazo_{etapa}",
                )
                responsavel_etapa = d2.selectbox(
                    f"Responsável — {etapa}",
                    list(opcoes_resp),
                    key=f"novo_responsavel_{etapa}",
                )
                etapas[etapa] = {
                    "data_planejada": prazo,
                    "responsavel_email": opcoes_resp[responsavel_etapa],
                }
        enviado = st.form_submit_button("Criar CN", type="primary")
    if enviado:
        dados = {
            "numero_cn": numero,
            "codigo_cn": codigo,
            "familia": familia,
            "descricao": descricao,
            "codigo_linha": codigo_linha,
            "responsavel_email": opcoes_resp[responsavel_label],
            "resumo_modificacao": resumo,
            "observacao": observacao,
            "criticidade": criticidade,
        }
        try:
            criar_cn(dados, etapas, usuario["email"])
        except (ValueError, SQLAlchemyError, PermissionError) as erro:
            logger.exception("Falha ao cadastrar CN.")
            st.error(f"Não foi possível cadastrar a CN: {erro}")
        else:
            st.success("CN criada. Uma notificação de confirmação foi enviada ao responsável.")
            st.rerun()


def _tela_importacao(usuario: dict[str, Any], responsaveis: pd.DataFrame) -> None:
    st.subheader("Importar planilha de CN's")
    arquivo = st.file_uploader(
        "Selecione a planilha Controle de CN's 2026 (.xlsm ou .xlsx)",
        type=["xlsm", "xlsx"],
        key="importar_planilha_cns",
    )
    if arquivo is None:
        st.info("Serão importados dados a partir da linha 3 da aba CONTROLE DE CN'S 2026.")
        return
    try:
        registros, alertas = analisar_planilha_cns(arquivo.getvalue())
    except (ValueError, OSError) as erro:
        st.error(str(erro))
        return
    try:
        salvar_referencias_cns(
            st.session_state["_referencias_cns"],
            usuario["email"],
        )
    except (PermissionError, SQLAlchemyError) as erro:
        logger.exception("Falha ao salvar listas de referência da planilha CN.")
        st.error(f"Não foi possível salvar as listas de Família e Descrição: {erro}")
        return
    st.success(
        "Cadastros mestres atualizados no banco e disponíveis para os usuários "
        "autorizados, mesmo sem importar as CNs."
    )
    if responsaveis.empty:
        st.warning("Cadastre pelo menos um responsável antes de importar.")
        return
    opcoes_resp = {
        f"{linha.usuario} · {linha.email} · {linha.setor}": linha.email
        for linha in responsaveis.itertuples(index=False)
    }
    st.caption(
        f"{len(registros)} registro(s) válido(s) encontrados. "
        "As datas serão tratadas exclusivamente como prazos; nenhuma etapa será marcada como concluída pela planilha."
    )
    if alertas:
        st.warning(f"{len(alertas)} célula(s) com data inválida ou múltiplas datas. Os detalhes serão registrados na auditoria.")
        st.dataframe(
            pd.DataFrame(alertas).rename(
                columns={"numero_cn": "CN", "coluna": "Coluna", "linha": "Linha", "detalhe": "Detalhe"}
            ),
            hide_index=True,
            use_container_width=True,
        )
    responsavel_label = st.selectbox(
        "Responsável inicial para CNs importadas",
        list(opcoes_resp),
        key="responsavel_importacao_cns",
    )
    if st.button("Importar CNs", type="primary", key="botao_importar_cns"):
        responsavel = opcoes_resp[responsavel_label]
        for registro in registros:
            registro["responsavel_email"] = responsavel
        try:
            resultado = importar_cns(registros, usuario["email"])
            for alerta in alertas:
                registrar_alerta_importacao(
                    usuario["email"],
                    alerta["numero_cn"],
                    alerta["coluna"],
                    alerta["linha"],
                    alerta["detalhe"],
                )
        except (ValueError, SQLAlchemyError, PermissionError) as erro:
            logger.exception("Falha ao importar planilha de CN's.")
            st.error(f"Não foi possível importar a planilha: {erro}")
        else:
            st.success(
                f"Importação concluída: {resultado['novas']} nova(s), "
                f"{resultado['duplicadas']} duplicada(s) ignorada(s). "
                f"{len(alertas)} alerta(s) de célula auditado(s)."
            )


def _tela_backup_cns(usuario: dict[str, Any]) -> None:
    st.subheader("Backup CN")
    st.caption(
        "Os backups são salvos no banco compartilhado e incluem CNs, etapas, "
        "confirmações e notificações. Usuários, responsáveis, cadastros mestres "
        "e auditoria não são removidos."
    )
    if st.button(
        "Gerar backup agora",
        type="primary",
        key="gerar_backup_cns",
    ):
        try:
            resultado = criar_backup_cns(usuario["email"])
        except (PermissionError, SQLAlchemyError) as erro:
            logger.exception("Falha ao gerar backup da base CN.")
            st.error(f"Não foi possível gerar o backup: {erro}")
        else:
            st.success(
                f"Backup #{resultado['id']} criado com "
                f"{resultado['quantidade_cns']} CN(s)."
            )
            st.download_button(
                "Baixar backup JSON",
                resultado["conteudo"],
                f"backup_cns_{resultado['id']}.json",
                "application/json",
                key=f"baixar_backup_cns_{resultado['id']}",
            )
    try:
        backups = listar_backups_cns(usuario["email"])
    except (PermissionError, SQLAlchemyError) as erro:
        logger.exception("Falha ao listar backups CN.")
        st.error(f"Não foi possível listar os backups: {erro}")
        return
    if backups.empty:
        st.info("Ainda não há backups de CN.")
        return
    st.dataframe(backups, hide_index=True, use_container_width=True)
    ids_backups = backups["id"].astype(int).tolist()
    backup_id = st.selectbox(
        "Backup para baixar",
        ids_backups,
        format_func=lambda item: f"Backup #{item}",
        key="selecionar_backup_cns",
    )
    try:
        conteudo = baixar_backup_cns(backup_id, usuario["email"])
    except (PermissionError, SQLAlchemyError, ValueError) as erro:
        logger.exception("Falha ao carregar backup CN para download.")
        st.error(f"Não foi possível carregar o backup: {erro}")
    else:
        st.download_button(
            "Baixar backup selecionado",
            conteudo,
            f"backup_cns_{backup_id}.json",
            "application/json",
            key=f"download_backup_selecionado_cns_{backup_id}",
        )


def _tela_limpar_dados_cns(usuario: dict[str, Any]) -> None:
    st.subheader("Limpar Dados CN")
    st.warning(
        "A limpeza gera um backup persistido antes de excluir CNs, etapas, "
        "confirmações e notificações. Usuários, permissões, responsáveis, "
        "cadastros mestres e auditoria serão preservados."
    )
    st.caption("Digite exatamente: LIMPAR DADOS CN")
    confirmacao = st.text_input(
        "Confirmação",
        key="confirmacao_limpeza_total_cns",
    )
    if st.button(
        "Gerar backup e limpar dados CN",
        type="secondary",
        disabled=confirmacao != "LIMPAR DADOS CN",
        key="limpar_dados_cns",
    ):
        try:
            resultado = limpar_dados_cns(usuario["email"])
            conteudo_backup = baixar_backup_cns(
                resultado["backup_id"],
                usuario["email"],
            )
        except (PermissionError, SQLAlchemyError, ValueError) as erro:
            logger.exception("Falha ao gerar backup e limpar a base CN.")
            st.error(f"Não foi possível concluir a limpeza: {erro}")
        else:
            st.success(
                f"Backup #{resultado['backup_id']} concluído. "
                f"{resultado['cns_removidas']} CN(s) removida(s); "
                "usuários, permissões e auditoria preservados."
            )
            st.download_button(
                "Baixar backup anterior à limpeza",
                conteudo_backup,
                f"backup_cns_{resultado['backup_id']}.json",
                "application/json",
                key=f"download_backup_limpeza_cns_{resultado['backup_id']}",
            )
            st.rerun()


def _tela_reimportar_base_cns(
    usuario: dict[str, Any],
    responsaveis: pd.DataFrame,
) -> None:
    st.subheader("Reimportar Base CN")
    st.caption(
        "Valida a planilha primeiro; em seguida, gera backup, substitui "
        "a base CN em uma transação e recalcula os indicadores no próximo carregamento."
    )
    arquivo = st.file_uploader(
        "Planilha base de CN's (.xlsm ou .xlsx)",
        type=["xlsm", "xlsx"],
        key="reimportar_base_planilha_cns",
    )
    if arquivo is None:
        st.info("A base atual permanecerá intacta até a confirmação da reimportação.")
        return
    try:
        registros, alertas = analisar_planilha_cns(arquivo.getvalue())
    except (ValueError, OSError) as erro:
        st.error(str(erro))
        return
    if responsaveis.empty:
        st.warning("Cadastre ao menos um responsável ativo antes de reimportar.")
        return
    opcoes_resp = {
        f"{linha.usuario} · {linha.email} · {linha.setor}": linha.email
        for linha in responsaveis.itertuples(index=False)
    }
    st.caption(f"{len(registros)} registro(s) válido(s) encontrados.")
    if alertas:
        st.warning(
            f"{len(alertas)} alerta(s) de data serão registrados após a importação."
        )
    responsavel_label = st.selectbox(
        "Responsável inicial para a nova base",
        list(opcoes_resp),
        key="responsavel_reimportacao_cns",
    )
    st.caption("Digite exatamente: REIMPORTAR BASE CN")
    confirmacao = st.text_input(
        "Confirmação",
        key="confirmacao_reimportacao_cns",
    )
    if st.button(
        "Gerar backup e reimportar base",
        type="primary",
        disabled=confirmacao != "REIMPORTAR BASE CN",
        key="confirmar_reimportacao_cns",
    ):
        responsavel = opcoes_resp[responsavel_label]
        for registro in registros:
            registro["responsavel_email"] = responsavel
        try:
            resultado = reimportar_base_cns(
                registros,
                usuario["email"],
                referencias=st.session_state["_referencias_cns"],
            )
            for alerta in alertas:
                registrar_alerta_importacao(
                    usuario["email"],
                    alerta["numero_cn"],
                    alerta["coluna"],
                    alerta["linha"],
                    alerta["detalhe"],
                )
        except (ValueError, SQLAlchemyError, PermissionError) as erro:
            logger.exception("Falha ao reimportar a base CN.")
            st.error(
                "A reimportação falhou. A transação manteve a base CN anterior: "
                f"{erro}"
            )
        else:
            st.success(
                f"Backup #{resultado['backup_id']} criado; "
                f"{resultado['cns_removidas']} CN(s) substituída(s) por "
                f"{resultado['novas']} nova(s), "
                f"{resultado['duplicadas']} duplicada(s) ignorada(s)."
            )
            st.rerun()


def _tela_acompanhamento(
    usuario: dict[str, Any],
    responsaveis: pd.DataFrame,
    dataframe: pd.DataFrame,
) -> None:
    detalhe_id = st.session_state.get("_cn_detalhe_id")
    if detalhe_id is not None:
        encontrados = dataframe[dataframe["id"] == detalhe_id]
        if not encontrados.empty:
            _render_detalhe_cn(usuario, responsaveis, encontrados.iloc[0])
            return
        st.session_state.pop("_cn_detalhe_id", None)

    filtrado = _filtrar_cns(dataframe)
    if filtrado.empty:
        st.info("Nenhuma CN encontrada com os filtros informados.")
        return
    exibicao = _dataframe_exibicao(filtrado).reset_index(drop=True)
    st.caption(f"{len(filtrado)} CN(s). Selecione uma linha para abrir os detalhes.")
    nonce = int(st.session_state.get("_tabela_cns_nonce", 0))
    selecao = st.dataframe(
        exibicao,
        hide_index=True,
        use_container_width=True,
        height=min(520, 42 + 38 * min(len(exibicao), 12)),
        on_select="rerun",
        selection_mode="single-row",
        key=f"tabela_acompanhamento_cns_{nonce}",
        column_config={
            "Número CN": st.column_config.TextColumn("CN", width="small"),
            "Descrição": st.column_config.TextColumn("Descrição", width="large"),
            "Família": st.column_config.TextColumn("Família", width="medium"),
            "Status": st.column_config.TextColumn("Status", width="medium"),
            "Responsável": st.column_config.TextColumn("Responsável", width="medium"),
            "Criticidade": st.column_config.TextColumn("Criticidade", width="small"),
        },
    )
    linhas_selecionadas = selecao.selection.rows
    if linhas_selecionadas:
        selecionado = filtrado.iloc[int(linhas_selecionadas[0])]
        st.session_state["_cn_detalhe_id"] = int(selecionado["id"])
        st.rerun()


def _render_detalhe_cn(
    usuario: dict[str, Any],
    responsaveis: pd.DataFrame,
    registro: pd.Series,
) -> None:
    mapa_nomes = st.session_state.get("_responsaveis_cns_nomes", {})
    botao_voltar, _ = st.columns([1, 5])
    if botao_voltar.button("← Voltar à lista", key="voltar_lista_detalhe_cn"):
        st.session_state.pop("_cn_detalhe_id", None)
        st.session_state["_tabela_cns_nonce"] = (
            int(st.session_state.get("_tabela_cns_nonce", 0)) + 1
        )
        st.rerun()

    numero = _formatar_identificador_cn(registro["numero_cn"])
    st.title(numero)
    st.subheader(str(registro["descricao"]))
    st.caption(f"Código interno: {registro['codigo_cn']}")
    st.caption(str(registro.get("resumo_modificacao") or ""))
    if registro.get("observacao"):
        with st.expander("Observação"):
            st.write(registro["observacao"])
    responsavel = mapa_nomes.get(
        registro["responsavel_email"], registro["responsavel_email"]
    )
    data_criacao = _formatar_data(registro["data_criacao"])
    data_criacao_rotulo = (
        data_criacao.strftime("%d/%m/%Y") if data_criacao else "—"
    )
    etapas = sorted(
        registro["etapas"] or [],
        key=lambda etapa: (
            ETAPAS_CN.index(etapa["etapa"])
            if etapa["etapa"] in ETAPAS_CN
            else len(ETAPAS_CN)
        ),
    )
    recebimentos_confirmados = sum(
        bool(etapa.get("data_recebimento")) for etapa in etapas
    )
    metadados = (
        ("Família", registro["familia"]),
        ("Responsável", responsavel),
        ("Data de criação", data_criacao_rotulo),
        ("Status geral", _normalizar_status_visual(registro["status"])),
        ("Criticidade", registro["criticidade"]),
        (
            "Recebimentos das etapas",
            f"{recebimentos_confirmados} de {len(etapas)} confirmados",
        ),
    )
    colunas = st.columns(len(metadados))
    for coluna, (titulo, valor) in zip(colunas, metadados):
        with coluna.container(border=True):
            st.caption(titulo)
            st.markdown(f"**{valor}**")

    concluidas = sum(
        etapa["status_calculado"] == "CONCLUÍDO" for etapa in etapas
    )
    atrasadas = sum(etapa["status_calculado"] == "ATRASADO" for etapa in etapas)
    pendentes = sum(etapa["status_calculado"] == "PENDENTE" for etapa in etapas)
    em_andamento = sum(
        etapa["status_calculado"] == "EM ANDAMENTO" for etapa in etapas
    )
    total_etapas = len(etapas)
    percentual = concluidas / total_etapas if total_etapas else 0.0
    st.subheader("Progresso das etapas")
    st.progress(
        percentual,
        text=f"{percentual:.0%} · {concluidas} de {total_etapas} etapas concluídas",
    )
    resumo_etapas = st.columns(4)
    for coluna, rotulo, valor in zip(
        resumo_etapas,
        ("Concluídas", "Pendentes", "Em andamento", "Atrasadas"),
        (concluidas, pendentes, em_andamento, atrasadas),
    ):
        coluna.metric(rotulo, valor)

    st.subheader("Etapas")
    responsaveis_por_email = {
        str(linha.email): str(linha.usuario)
        for linha in responsaveis.itertuples(index=False)
    }
    colunas_etapas = st.columns(2)
    hoje = date.today()
    for indice, etapa in enumerate(etapas):
        coluna = colunas_etapas[indice % 2]
        with coluna.container(border=True):
            status_etapa = etapa["status_calculado"]
            prazo = _formatar_data(etapa.get("data_planejada"))
            conclusao = _formatar_data(etapa.get("data_conclusao"))
            if status_etapa == "CONCLUÍDO":
                icone = "✅"
                rotulo_status = "Concluído"
            elif status_etapa == "ATRASADO":
                icone = "🔴"
                rotulo_status = "Atrasado"
            elif status_etapa == "PENDENTE":
                icone = "⚪"
                rotulo_status = "Não iniciado"
            elif prazo is not None and 0 <= (prazo - hoje).days <= 7:
                icone = "🟡"
                rotulo_status = "Próximo do vencimento"
            else:
                icone = "🔵"
                rotulo_status = "Em andamento"
            st.markdown(f"**{icone} {etapa['etapa']}**")
            st.caption(rotulo_status)
            st.write(f"Prazo: {prazo.strftime('%d/%m/%Y') if prazo else 'Não definido'}")
            if conclusao:
                st.write(f"Conclusão: {conclusao.strftime('%d/%m/%Y')}")
            if status_etapa == "ATRASADO" and prazo:
                st.error(f"Atrasado há {(hoje - prazo).days} dia(s)")
            elif status_etapa == "EM ANDAMENTO" and prazo is not None:
                dias_restantes = (prazo - hoje).days
                if 0 <= dias_restantes <= 7:
                    st.warning(
                        "Vence hoje"
                        if dias_restantes == 0
                        else f"Vence em {dias_restantes} dia(s)"
                    )
            st.write(
                "Responsável: "
                + responsaveis_por_email.get(
                    str(etapa["responsavel_email"]),
                    str(etapa["responsavel_email"]),
                )
            )
            recebida_em = _formatar_data(etapa.get("data_recebimento"))
            if recebida_em:
                nome_recebedor = responsaveis_por_email.get(
                    str(etapa.get("recebido_por") or ""),
                    str(etapa.get("recebido_por") or ""),
                )
                st.success(
                    f"Recebimento confirmado em {recebida_em:%d/%m/%Y}"
                    + (f" por {nome_recebedor}" if nome_recebedor else "")
                )
            else:
                responsavel_etapa = str(etapa["responsavel_email"])
                pode_confirmar = (
                    responsavel_etapa.casefold() == usuario["email"].casefold()
                    or usuario["perfil"] in PAPEIS_AUDITORIA
                )
                if pode_confirmar and st.button(
                    "☑ Confirmar recebimento",
                    type="primary",
                    key=f"confirmar_etapa_{registro['id']}_{etapa['id']}",
                    use_container_width=True,
                ):
                    try:
                        confirmar_recebimento_etapa(
                            int(registro["id"]),
                            etapa["etapa"],
                            usuario["email"],
                        )
                    except (ValueError, PermissionError, SQLAlchemyError) as erro:
                        logger.exception(
                            "Falha ao confirmar recebimento da etapa da CN."
                        )
                        st.error(f"Não foi possível confirmar o recebimento: {erro}")
                    else:
                        st.rerun()
                else:
                    st.caption("Recebimento pendente do responsável da etapa.")

    if responsaveis.empty:
        return
    labels = {
        f"{linha.usuario} · {linha.email}": linha.email
        for linha in responsaveis.itertuples(index=False)
    }
    emails_responsaveis = list(labels.values())
    with st.expander("Atualizar etapa"):
        etapa_selecionada = st.selectbox(
            "Etapa",
            ETAPAS_CN,
            key=f"etapa_editar_{registro['id']}",
        )
        etapa_atual = next(
            (etapa for etapa in etapas if etapa["etapa"] == etapa_selecionada), {}
        )
        prazo_atual = _formatar_data(etapa_atual.get("data_planejada"))
        conclusao_atual = _formatar_data(etapa_atual.get("data_conclusao"))
        atual_email = str(etapa_atual.get("responsavel_email", ""))
        idx_email = (
            emails_responsaveis.index(atual_email)
            if atual_email in emails_responsaveis
            else 0
        )
        pode_editar = (
            usuario["perfil"] in PAPEIS_AUDITORIA
            or atual_email.casefold() == usuario["email"].casefold()
        )
        if pode_editar:
            with st.form(f"editar_etapa_cn_{registro['id']}"):
                p1, p2, p3 = st.columns(3)
                prazo_novo = p1.date_input(
                    "Data planejada (prazo)",
                    value=prazo_atual,
                    key=f"prazo_editar_{registro['id']}_{etapa_selecionada}",
                )
                conclusao_nova = p2.date_input(
                    "Data real de conclusão",
                    value=conclusao_atual,
                    key=f"conclusao_editar_{registro['id']}_{etapa_selecionada}",
                )
                responsavel_novo = p3.selectbox(
                    "Responsável pela etapa",
                    emails_responsaveis,
                    index=idx_email,
                    format_func=lambda email: next(
                        (nome for nome, valor in labels.items() if valor == email),
                        email,
                    ),
                    key=f"responsavel_editar_{registro['id']}_{etapa_selecionada}",
                )
                salvar_etapa = st.form_submit_button("Salvar etapa")
            if salvar_etapa:
                try:
                    atualizar_etapa(
                        int(registro["id"]),
                        etapa_selecionada,
                        prazo_novo,
                        conclusao_nova,
                        responsavel_novo,
                        usuario["email"],
                    )
                except (ValueError, SQLAlchemyError, PermissionError) as erro:
                    logger.exception("Falha ao atualizar etapa da CN.")
                    st.error(f"Não foi possível atualizar a etapa: {erro}")
                else:
                    st.success("Etapa atualizada.")
                    st.rerun()
        else:
            st.info(
                "Somente o responsável designado ou os perfis ADMIN/QUALIDADE "
                "podem editar esta etapa."
            )

    if usuario["perfil"] in PAPEIS_AUDITORIA:
        with st.expander("Editar dados gerais da CN"):
            responsavel_map = {
                f"{linha.usuario} · {linha.email}": linha.email
                for linha in responsaveis.itertuples(index=False)
            }
            rotulos_resp = list(responsavel_map)
            atual_resp = str(registro["responsavel_email"])
            atual_idx = next(
                (
                    index
                    for index, label in enumerate(rotulos_resp)
                    if responsavel_map[label] == atual_resp
                ),
                0,
            )
            with st.form(f"editar_dados_cn_{registro['id']}"):
                familia = st.text_input("Família", value=registro["familia"])
                descricao = st.text_input("Descrição", value=registro["descricao"])
                st.text_input(
                    "Código Linha",
                    value=registro["codigo_linha"],
                    disabled=True,
                )
                codigo_cn = st.text_input(
                    "Código interno", value=registro["codigo_cn"]
                )
                resumo = st.text_area(
                    "Resumo da Modificação",
                    value=registro["resumo_modificacao"],
                )
                observacao = st.text_area(
                    "Observação", value=registro["observacao"]
                )
                responsavel_label = st.selectbox(
                    "Responsável",
                    rotulos_resp,
                    index=atual_idx,
                    key=f"responsavel_geral_cn_{registro['id']}",
                )
                criticidade = st.selectbox(
                    "Criticidade",
                    CRITICIDADES,
                    index=CRITICIDADES.index(registro["criticidade"]),
                    key=f"criticidade_cn_{registro['id']}",
                )
                salvar_dados = st.form_submit_button("Salvar dados da CN")
            if salvar_dados:
                try:
                    referencias = carregar_referencias_cns()
                    linha_automatica = next(
                        (
                            linha["codigo_linha"]
                            for linha in referencias["linhas"]
                            if linha["familia"].casefold() == familia.casefold()
                            and linha["descricao"].casefold()
                            == descricao.casefold()
                        ),
                        registro["codigo_linha"],
                    )
                    atualizar_cn(
                        int(registro["id"]),
                        {
                            "familia": familia,
                            "descricao": descricao,
                            "codigo_linha": linha_automatica,
                            "codigo_cn": codigo_cn,
                            "resumo_modificacao": resumo,
                            "observacao": observacao,
                            "responsavel_email": responsavel_map[responsavel_label],
                            "criticidade": criticidade,
                        },
                        usuario["email"],
                    )
                except (ValueError, SQLAlchemyError, PermissionError) as erro:
                    logger.exception("Falha ao editar dados gerais da CN.")
                    st.error(f"Não foi possível atualizar a CN: {erro}")
                else:
                    st.success("Dados da CN atualizados.")
                    st.rerun()


def _tela_auditoria(
    usuario: dict[str, Any],
    somente_administrador: bool = False,
) -> None:
    if somente_administrador and usuario["perfil"] != "ADMIN":
        st.error("Somente o perfil ADMIN pode acessar a Administração CN's.")
        return
    if not somente_administrador and usuario["perfil"] not in PAPEIS_AUDITORIA:
        st.error("Seu perfil não tem permissão para acessar a Auditoria CN.")
        return
    st.subheader("📜 Auditoria CN")
    try:
        auditoria = (
            carregar_auditoria_administracao(usuario["email"])
            if somente_administrador
            else carregar_auditoria()
        )
    except SQLAlchemyError as erro:
        logger.exception("Falha ao consultar auditoria das CNs.")
        st.error(f"Não foi possível carregar a auditoria: {erro}")
        return
    if not auditoria.empty:
        filtros = st.columns(4)
        for controle, campo in zip(filtros, ("CN", "Usuário", "Campo", "Tipo Alteração")):
            valores = sorted(auditoria[campo].dropna().astype(str).unique())
            selecionado = controle.selectbox(
                campo,
                ["Todos", *valores],
                key=f"filtro_auditoria_cn_{campo}",
            )
            if selecionado != "Todos":
                auditoria = auditoria[auditoria[campo].astype(str) == selecionado]
        periodo = st.date_input(
            "Período",
            value=(date.today().replace(day=1), date.today()),
            key="periodo_auditoria_cns",
        )
        if isinstance(periodo, tuple) and len(periodo) == 2:
            datas = pd.to_datetime(auditoria["Data/Hora"], errors="coerce")
            auditoria = auditoria[
                datas.dt.date.between(periodo[0], periodo[1])
            ]
    st.dataframe(auditoria, hide_index=True, use_container_width=True)
    e1, e2 = st.columns(2)
    e1.download_button(
        "Exportar Excel",
        exportar_dataframe_excel(auditoria, "Auditoria CN"),
        "auditoria_cns.xlsx",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        key="exportar_auditoria_cns_excel",
    )
    e2.download_button(
        "Exportar CSV",
        exportar_dataframe_csv(auditoria),
        "auditoria_cns.csv",
        "text/csv",
        key="exportar_auditoria_cns_csv",
    )


def _tela_notificacoes(usuario: dict[str, Any]) -> None:
    st.subheader("Central de notificações")
    try:
        notificacoes = carregar_notificacoes(usuario["email"])
    except SQLAlchemyError as erro:
        logger.exception("Falha ao carregar notificações das CNs.")
        st.error(f"Não foi possível carregar as notificações: {erro}")
        return
    if notificacoes.empty:
        st.info("Não há notificações para este usuário.")
        return
    st.dataframe(
        notificacoes[["Tipo", "CN", "Notificação", "Data/Hora", "Lida"]],
        hide_index=True,
        use_container_width=True,
    )
    nao_lidas = notificacoes[notificacoes["Lida"] == 0]
    if not nao_lidas.empty:
        opcoes = {
            f"{linha.Tipo} · CN {linha.CN} · {linha.Notificação}": int(linha.id)
            for linha in nao_lidas.itertuples(index=False)
        }
        escolhido = st.selectbox("Marcar como lida", list(opcoes), key="notificacao_cns_lida")
        if st.button("Marcar notificação como lida", key="marcar_notificacao_cns_lida"):
            try:
                marcar_notificacao_lida(opcoes[escolhido], usuario["email"])
            except SQLAlchemyError as erro:
                logger.exception("Falha ao marcar notificação como lida.")
                st.error(f"Não foi possível atualizar a notificação: {erro}")
            else:
                st.rerun()


def _tela_responsaveis_administracao(usuario: dict[str, Any]) -> None:
    st.subheader("Gestão de responsáveis")
    try:
        responsaveis = listar_responsaveis_administracao(usuario["email"])
        usuarios = listar_usuarios("ATIVO")
    except SQLAlchemyError as erro:
        logger.exception("Falha ao carregar administração de responsáveis CN.")
        st.error(f"Não foi possível carregar os responsáveis: {erro}")
        return

    opcoes = {"Novo responsável": None}
    for linha in responsaveis.itertuples(index=False):
        estado = "Ativo" if linha.ativo else "Inativo"
        opcoes[f"{linha.usuario} · {linha.email} · {estado}"] = int(linha.id)
    escolha = st.selectbox(
        "Responsável",
        list(opcoes),
        key="admin_selecao_responsavel_cn",
    )
    responsavel_id = opcoes[escolha]
    atual = (
        responsaveis[responsaveis["id"] == responsavel_id].iloc[0]
        if responsavel_id is not None
        else None
    )

    if usuarios.empty and atual is None:
        st.info("Cadastre um usuário ativo antes de adicioná-lo como responsável.")
        return

    with st.form(f"form_admin_responsavel_cn_{responsavel_id or 'novo'}"):
        nome = st.text_input(
            "Usuário",
            value=str(atual["usuario"]) if atual is not None else "",
        )
        email = st.text_input(
            "E-mail",
            value=str(atual["email"]) if atual is not None else "",
        )
        setor = st.text_input(
            "Setor",
            value=str(atual["setor"]) if atual is not None else "",
        )
        ativo = st.checkbox(
            "Responsável ativo",
            value=bool(atual["ativo"]) if atual is not None else True,
        )
        salvar = st.form_submit_button("Salvar responsável", type="primary")
    if salvar:
        try:
            administrar_responsavel(
                responsavel_id,
                nome,
                email,
                setor,
                ativo,
                usuario["email"],
            )
        except (ValueError, PermissionError, SQLAlchemyError) as erro:
            logger.exception("Falha ao salvar responsável pela Administração CN's.")
            st.error(f"Não foi possível salvar o responsável: {erro}")
        else:
            st.success("Responsável atualizado.")
            st.rerun()

    st.dataframe(
        responsaveis.rename(
            columns={
                "usuario": "Usuário",
                "email": "E-mail",
                "setor": "Setor",
                "ativo": "Ativo",
            }
        ),
        hide_index=True,
        use_container_width=True,
    )


def _tela_cadastros_mestres(usuario: dict[str, Any]) -> None:
    st.subheader("Famílias, descrições e código da linha")
    try:
        referencias = carregar_referencias_cns()
    except SQLAlchemyError as erro:
        logger.exception("Falha ao carregar os cadastros mestres de CN.")
        st.error(f"Não foi possível carregar os cadastros mestres: {erro}")
        return

    if usuario["perfil"] in PAPEIS_AUDITORIA:
        st.caption(
            "Os cadastros são compartilhados pelo banco de dados. Edite a tabela "
            "ou adicione linhas e salve; a planilha serve para importação inicial "
            "e atualização em massa."
        )
        tabela = pd.DataFrame(
            referencias["linhas"],
            columns=["familia", "descricao", "codigo_linha"],
        )
        chave_tabela = hash(
            tuple(
                (linha["familia"], linha["descricao"], linha["codigo_linha"])
                for linha in referencias["linhas"]
            )
        )
        with st.form("form_cadastros_mestres_cns"):
            editada = st.data_editor(
                tabela,
                num_rows="dynamic",
                hide_index=True,
                use_container_width=True,
                column_config={
                    "familia": st.column_config.TextColumn("Família"),
                    "descricao": st.column_config.TextColumn("Descrição"),
                    "codigo_linha": st.column_config.TextColumn("Código da Linha"),
                },
                key=f"tabela_cadastros_mestres_cns_{chave_tabela}",
            )
            salvar = st.form_submit_button(
                "Salvar cadastros mestres",
                type="primary",
            )
        if salvar:
            try:
                salvar_referencias_cns(
                    {"linhas": editada.to_dict("records")},
                    usuario["email"],
                )
            except (PermissionError, SQLAlchemyError, ValueError) as erro:
                logger.exception("Falha ao salvar os cadastros mestres de CN.")
                st.error(f"Não foi possível salvar os cadastros mestres: {erro}")
            else:
                st.success("Cadastros mestres salvos no banco de dados.")
                st.rerun()
    else:
        st.info("Seu perfil tem acesso de consulta aos cadastros mestres.")
        if referencias["linhas"]:
            st.dataframe(
                pd.DataFrame(referencias["linhas"]).rename(
                    columns={
                        "familia": "Família",
                        "descricao": "Descrição",
                        "codigo_linha": "Código da Linha",
                    }
                ),
                hide_index=True,
                use_container_width=True,
            )
        else:
            st.info("Ainda não há famílias, descrições ou códigos de linha cadastrados.")


def render_administracao_cadastros_mestres(
    usuario: dict[str, Any] | None = None,
) -> None:
    usuario = usuario or _render_autenticacao("Administração")
    if usuario is None:
        return
    try:
        usuario_atual = usuario_ativo(usuario["email"])
    except SQLAlchemyError as erro:
        logger.exception("Falha ao validar acesso aos cadastros mestres.")
        st.error(f"Não foi possível validar sua sessão: {erro}")
        return
    if usuario_atual is None or usuario_atual["perfil"] != "ADMIN":
        st.error("Somente o perfil ADMIN pode acessar os cadastros mestres.")
        return

    if usuario["perfil"] == "ADMIN":
        mestres, responsaveis = st.tabs(
            ["Famílias, descrições e código da linha", "Responsáveis"]
        )
        with mestres:
            _tela_cadastros_mestres(usuario)
        with responsaveis:
            _tela_responsaveis_administracao(usuario)
    else:
        _tela_cadastros_mestres(usuario)
        st.subheader("Responsáveis ativos")
        try:
            responsaveis = carregar_responsaveis()
        except SQLAlchemyError as erro:
            logger.exception("Falha ao consultar os responsáveis cadastrados.")
            st.error(f"Não foi possível carregar os responsáveis: {erro}")
        else:
            st.dataframe(
                responsaveis.rename(
                    columns={
                        "usuario": "Usuário",
                        "email": "E-mail",
                        "setor": "Setor",
                    }
                ),
                hide_index=True,
                use_container_width=True,
            )


def _tela_notificacoes_administracao(usuario: dict[str, Any]) -> None:
    st.subheader("Central de notificações CN")
    try:
        notificacoes = carregar_todas_notificacoes_administracao(usuario["email"])
    except (PermissionError, SQLAlchemyError) as erro:
        logger.exception("Falha ao carregar notificações na Administração CN's.")
        st.error(f"Não foi possível carregar as notificações: {erro}")
        return
    if notificacoes.empty:
        st.info("Não há notificações registradas.")
        return
    st.dataframe(notificacoes, hide_index=True, use_container_width=True)


def _tela_limpeza_importacoes(
    usuario: dict[str, Any],
    dataframe: pd.DataFrame,
) -> None:
    st.subheader("Limpar CNs importadas")
    st.warning(
        "Esta ação remove permanentemente apenas CNs identificadas como criadas "
        "pela importação da planilha, incluindo etapas, confirmações e notificações. "
        "Os registros de auditoria serão preservados."
    )
    try:
        numeros_importados = set(
            listar_numeros_cns_importadas(usuario["email"])
        )
    except (PermissionError, SQLAlchemyError) as erro:
        logger.exception("Falha ao identificar CNs importadas para limpeza.")
        st.error(f"Não foi possível identificar as CNs importadas: {erro}")
        return

    registros_importados = dataframe[
        dataframe["numero_cn"].astype(str).isin(numeros_importados)
    ]
    st.metric("CNs importadas que serão removidas", len(registros_importados))
    if registros_importados.empty:
        st.info("Não há CNs importadas ativas para limpar.")
        return

    st.dataframe(
        _dataframe_exibicao(registros_importados),
        hide_index=True,
        use_container_width=True,
    )
    st.caption("Para habilitar a limpeza, digite exatamente: LIMPAR IMPORTAÇÕES")
    confirmacao = st.text_input(
        "Confirmação",
        key="confirmacao_limpeza_importacoes_cns",
    )
    if st.button(
        "Limpar CNs importadas",
        type="secondary",
        disabled=confirmacao != "LIMPAR IMPORTAÇÕES",
        key="limpar_importacoes_cns",
    ):
        try:
            quantidade = limpar_cns_importadas(usuario["email"])
        except (PermissionError, SQLAlchemyError, ValueError) as erro:
            logger.exception("Falha ao limpar CNs importadas.")
            st.error(f"Não foi possível limpar as CNs importadas: {erro}")
        else:
            st.session_state["_resultado_limpeza_importacoes_cns"] = (
                f"{quantidade} CN(s) importada(s) removida(s). "
                "A auditoria foi preservada."
            )
            st.rerun()


def _tela_registros_administracao(
    usuario: dict[str, Any],
    responsaveis: pd.DataFrame,
    dataframe: pd.DataFrame,
) -> None:
    st.subheader("Gerenciar registros CN")
    if dataframe.empty:
        st.info("Não há CNs cadastradas.")
        return

    selecionado_id = st.session_state.get("_cn_admin_detalhe_id")
    selecionado = dataframe[dataframe["id"] == selecionado_id]
    if selecionado.empty:
        st.session_state.pop("_cn_admin_detalhe_id", None)
        filtrado = _filtrar_cns(dataframe)
        if filtrado.empty:
            st.info("Nenhuma CN encontrada com os filtros informados.")
            return
        st.caption(f"{len(filtrado)} CN(s). Selecione uma linha para abrir os detalhes.")
        nonce = int(st.session_state.get("_tabela_admin_cns_nonce", 0))
        selecao = st.dataframe(
            _dataframe_exibicao(filtrado).reset_index(drop=True),
            hide_index=True,
            use_container_width=True,
            height=min(520, 42 + 38 * min(len(filtrado), 12)),
            on_select="rerun",
            selection_mode="single-row",
            key=f"tabela_admin_cns_{nonce}",
            column_config={
                "Número CN": st.column_config.TextColumn("CN", width="small"),
                "Descrição": st.column_config.TextColumn("Descrição", width="large"),
                "Família": st.column_config.TextColumn("Família", width="medium"),
                "Status": st.column_config.TextColumn("Status", width="small"),
                "Responsável": st.column_config.TextColumn(
                    "Responsável", width="medium"
                ),
                "Criticidade": st.column_config.TextColumn(
                    "Criticidade", width="small"
                ),
            },
        )
        linhas_selecionadas = selecao.selection.rows
        if linhas_selecionadas:
            registro = filtrado.iloc[int(linhas_selecionadas[0])]
            st.session_state["_cn_admin_detalhe_id"] = int(registro["id"])
            st.rerun()
        return

    registro = selecionado.iloc[0]
    st.divider()
    st.warning(
        "A exclusão remove a CN, suas etapas, confirmações e notificações. "
        "O evento de exclusão permanecerá na auditoria."
    )
    confirmar_exclusao = st.checkbox(
        f"Confirmo a exclusão definitiva de {_formatar_identificador_cn(registro['numero_cn'])}",
        key=f"confirmar_excluir_cn_{registro['id']}",
    )
    if st.button(
        "Excluir CN",
        type="secondary",
        disabled=not confirmar_exclusao,
        key=f"excluir_cn_admin_{registro['id']}",
    ):
        try:
            excluir_cn(int(registro["id"]), usuario["email"])
        except (ValueError, PermissionError, SQLAlchemyError) as erro:
            logger.exception("Falha ao excluir CN pela administração.")
            st.error(f"Não foi possível excluir a CN: {erro}")
        else:
            st.session_state.pop("_cn_admin_detalhe_id", None)
            st.success("CN excluída; o evento foi preservado na auditoria.")
            st.rerun()
    _render_detalhe_cn(usuario, responsaveis, registro)


def _tela_administracao_cns(
    usuario: dict[str, Any],
    responsaveis: pd.DataFrame,
    dataframe: pd.DataFrame,
) -> None:
    if usuario["perfil"] != "ADMIN":
        st.error("Somente o perfil ADMIN pode acessar a Administração CN's.")
        return

    st.subheader("⚙ Administração CN's")
    menu = st.radio(
        "Seção administrativa",
        [
            "Registros CN",
            "Limpar importações",
            "📦 Backup CN",
            "🔄 Reimportar Base",
            "🗑 Limpar Dados CN",
            "Notificações CN",
            "Auditoria CN",
        ],
        horizontal=True,
        key="menu_administracao_cns",
        label_visibility="collapsed",
    )
    if menu == "Registros CN":
        _tela_registros_administracao(usuario, responsaveis, dataframe)
    elif menu == "Limpar importações":
        _tela_limpeza_importacoes(usuario, dataframe)
    elif menu == "📦 Backup CN":
        _tela_backup_cns(usuario)
    elif menu == "🔄 Reimportar Base":
        _tela_reimportar_base_cns(usuario, responsaveis)
    elif menu == "🗑 Limpar Dados CN":
        _tela_limpar_dados_cns(usuario)
    elif menu == "Notificações CN":
        _tela_notificacoes_administracao(usuario)
    elif menu == "Auditoria CN":
        _tela_auditoria(usuario, somente_administrador=True)


@st.fragment(run_every="60s")
def _processar_escalonamentos_periodicamente() -> None:
    try:
        notificacoes_criadas = processar_notificacoes_cns()
    except SQLAlchemyError as erro:
        logger.exception("Falha ao processar escalonamentos de CN.")
        st.error(f"Não foi possível processar os lembretes das CNs: {erro}")
    else:
        if notificacoes_criadas:
            st.toast(
                f"{notificacoes_criadas} nova(s) notificação(ões) de CN disponível(is)."
            )


def _garantir_responsavel_usuario(
    usuario: dict[str, Any],
    database_url: str | None = None,
) -> pd.DataFrame:
    responsaveis = carregar_responsaveis(database_url)
    if usuario["perfil"] not in PAPEIS_AUDITORIA:
        return responsaveis

    emails_existentes = {
        str(email).strip().casefold() for email in responsaveis["email"]
    }
    usuarios_ativos = listar_usuarios("ATIVO")
    novos_responsaveis = usuarios_ativos[
        usuarios_ativos["perfil"].isin(PAPEIS_AUDITORIA)
        & ~usuarios_ativos["email"].str.casefold().isin(emails_existentes)
    ]
    for candidato in novos_responsaveis.itertuples(index=False):
        salvar_responsavel(
            str(candidato.nome),
            str(candidato.email),
            "Não informado",
            usuario["email"],
            database_url,
        )
    if not novos_responsaveis.empty:
        responsaveis = carregar_responsaveis(database_url)
    return responsaveis


def render_gestao_cns() -> None:
    st.markdown("## 📋 Gestão de CN's")
    usuario = _render_autenticacao()
    if usuario is None:
        return
    resultado_limpeza = st.session_state.pop(
        "_resultado_limpeza_importacoes_cns",
        None,
    )
    if resultado_limpeza:
        st.success(resultado_limpeza)
    topo, sair = st.columns([4, 1])
    topo.caption(f"Sessão: {usuario['nome']} · Perfil {usuario['perfil']}")
    if sair.button("Sair", key="sair_gestao_cns", use_container_width=True):
        st.session_state.pop("_usuario_gestao_cns", None)
        st.rerun()
    try:
        inicializar_banco()
        dataframe = carregar_cns()
        responsaveis = _garantir_responsavel_usuario(usuario)
        referencias = carregar_referencias_cns()
    except SQLAlchemyError as erro:
        logger.exception("Falha ao inicializar a Gestão de CN's.")
        st.error(f"Não foi possível carregar os dados das CNs: {erro}")
        return
    if _engine().dialect.name == "sqlite":
        st.warning(
            "O banco configurado é SQLite local. No Streamlit Cloud, os dados "
            "podem desaparecer após reinicializações ou novos deploys. Configure "
            "um banco PostgreSQL persistente em Manage app → Settings → Secrets, "
            "usando a chave QMAX_CALIBRACAO_DATABASE_URL."
        )
    st.session_state["_responsaveis_cns_nomes"] = {
        str(linha.email): str(linha.usuario)
        for linha in responsaveis.itertuples(index=False)
    }
    _processar_escalonamentos_periodicamente()
    opcoes = [
        "Dashboard",
        "Acompanhamento",
        "Gantt anual",
        "Notificações",
    ]
    if usuario["perfil"] in PAPEIS_AUDITORIA:
        opcoes.extend(["Cadastrar CN", "Importar planilha", "📜 Auditoria CN"])
    if usuario["perfil"] == "ADMIN":
        opcoes.append("⚙ Administração CN's")
    menu = st.radio(
        "Menu Gestão de CN's",
        opcoes,
        horizontal=True,
        key="menu_gestao_cns",
        label_visibility="collapsed",
    )
    if menu == "Dashboard":
        _gerar_dashboard(dataframe)
    elif menu == "Acompanhamento":
        _tela_acompanhamento(usuario, responsaveis, dataframe)
    elif menu == "Gantt anual":
        _render_gantt_anual(dataframe)
    elif menu == "Cadastrar CN":
        _tela_cadastro(usuario, responsaveis, referencias)
    elif menu == "Importar planilha":
        _tela_importacao(usuario, responsaveis)
    elif menu == "Notificações":
        _tela_notificacoes(usuario)
    elif menu == "📜 Auditoria CN":
        _tela_auditoria(usuario)
    elif menu == "⚙ Administração CN's":
        _tela_administracao_cns(usuario, responsaveis, dataframe)


__all__ = [
    "analisar_planilha_cns",
    "render_administracao_cadastros_mestres",
    "render_gestao_cns",
]
