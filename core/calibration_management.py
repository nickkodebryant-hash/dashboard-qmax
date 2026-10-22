from __future__ import annotations

from datetime import date, datetime, timedelta
from io import BytesIO
import hashlib
import logging
import re
import unicodedata
import warnings
from typing import Any, Iterable

import openpyxl
import pandas as pd
import plotly.express as px
import plotly.graph_objects as go
import streamlit as st
from openpyxl.utils.datetime import from_excel
from pypdf import PdfWriter
from pypdf.generic import (
    DecodedStreamObject,
    DictionaryObject,
    NameObject,
)
from sqlalchemy.exc import SQLAlchemyError

from core.authentication import render_autenticacao
from database.calibration_store import (
    PAPEIS_AUDITORIA,
    PERFIS_USUARIO,
    alterar_perfil_usuario,
    aprovar_solicitacao,
    autenticar_usuario,
    carregar_auditoria_alteracoes,
    carregar_auditoria_eventos,
    carregar_equipamentos,
    carregar_historico_equipamento,
    criar_usuario,
    exportar_dataframe_csv,
    exportar_dataframe_excel,
    importar_equipamentos,
    listar_usuarios,
    normalizar_codigo,
    rejeitar_solicitacao,
    usuario_ativo,
)


logger = logging.getLogger(__name__)

COLUNAS_CALIBRACAO = {
    "NUMERO_EQUIPAMENTO": (
        "Nº EQUIPAMENTO",
        "NUMERO EQUIPAMENTO",
        "N EQUIPAMENTO",
        "NO EQUIPAMENTO",
    ),
    "DESCRICAO": ("DESCRIÇÃO", "DESCRICAO", "EQUIPAMENTO"),
    "UNIDADE": ("UNIDADE",),
    "CAPACIDADE": ("CAPACIDADE", "CAPACIDADE INSTRUMENTO"),
    "VALOR_DIVISAO_RESOLUCAO": (
        "VALOR DIVISÃO / RESOLUÇÃO",
        "VALOR DIVISAO RESOLUCAO",
        "VALOR DIVISÃO (RESOLUÇÃO)",
        "VALOR DIVISAO RESOLUCAO",
    ),
    "FAIXA_USO": ("FAIXA DE USO", "FAIXA USO"),
    "LOCAL_USO": ("LOCAL DE USO", "LOCAL USO"),
    "PROXIMA_CALIBRACAO": ("PRÓXIMA CALIBRAÇÃO", "PROXIMA CALIBRACAO"),
    "STATUS_CALIBRACAO": ("STATUS CALIBRAÇÃO", "STATUS CALIBRACAO"),
    "DATA_CALIBRACAO": ("DATA CALIBRAÇÃO", "DATA CALIBRACAO"),
    "SETOR": ("SETOR",),
    "FABRICANTE": ("FABRICANTE",),
    "STATUS_EQUIPAMENTO": ("STATUS EQUIPAMENTO",),
    "FREQUENCIA_CALIBRACAO": (
        "FREQUÊNCIA CALIBRAÇÃO",
        "FREQUENCIA CALIBRACAO",
        "VALIDADE CALIBRAÇÃO",
        "VALIDADE CALIBRACAO",
    ),
}
COLUNAS_OBRIGATORIAS_CAL = (
    "NUMERO_EQUIPAMENTO",
    "DESCRICAO",
    "PROXIMA_CALIBRACAO",
    "STATUS_CALIBRACAO",
)
CAMPOS_RELATORIO_CAL = (
    "Nº Equipamento",
    "Descrição",
    "Unidade",
    "Capacidade",
    "Local de Uso",
    "Próxima Calibração",
    "Status Calibração",
    "Fabricante",
    "Setor",
    "Status Equipamento",
)
ROTULOS_COLUNAS = {
    "NUMERO_EQUIPAMENTO": "Nº Equipamento",
    "DESCRICAO": "Descrição",
    "UNIDADE": "Unidade",
    "CAPACIDADE": "Capacidade",
    "VALOR_DIVISAO_RESOLUCAO": "Valor Divisão / Resolução",
    "FAIXA_USO": "Faixa de Uso",
    "LOCAL_USO": "Local de Uso",
    "PROXIMA_CALIBRACAO": "Próxima Calibração",
    "STATUS_CALIBRACAO": "Status Calibração",
    "DATA_CALIBRACAO": "Data Calibração",
    "SETOR": "Setor",
    "FABRICANTE": "Fabricante",
    "STATUS_EQUIPAMENTO": "Status Equipamento",
    "DIAS_RESTANTES": "Dias Restantes",
}
STATUS_CALIBRACAO = ("VENCIDO", "A VENCER", "NO PRAZO", "SEM PRÓXIMA CALIBRAÇÃO")
STATUS_INATIVOS = {"RETIRADO DE USO", "EXTRAVIADO"}
CORES_STATUS_CALIBRACAO = {
    "VENCIDO": "#e74c3c",
    "A VENCER": "#f39c12",
    "NO PRAZO": "#2ecc71",
    "SEM PRÓXIMA CALIBRAÇÃO": "#95a5a6",
}
PALETA_GRAFICOS_CALIBRACAO = (
    "#e99a2b",
    "#4da3ff",
    "#2ecc71",
    "#b07ce8",
    "#f06f68",
    "#32c5c0",
    "#f4cf58",
    "#7189ff",
    "#ed7db7",
    "#9abf45",
)


class ErroEstruturaCalibracao(ValueError):
    pass


def _normalizar_cabecalho(valor: Any) -> str:
    if valor is None:
        return ""
    texto = unicodedata.normalize("NFKD", str(valor))
    texto = "".join(caractere for caractere in texto if not unicodedata.combining(caractere))
    texto = texto.replace("º", " ").replace("°", " ")
    return " ".join(re.sub(r"[^A-Z0-9]+", " ", texto.upper()).split())


def _indice_cabecalhos(valores: Iterable[Any]) -> dict[str, int]:
    indices: dict[str, int] = {}
    for indice, valor in enumerate(valores):
        chave = _normalizar_cabecalho(valor)
        if chave:
            indices.setdefault(chave, indice)
    return indices


def _mapear_colunas(valores: Iterable[Any]) -> dict[str, int]:
    indices = _indice_cabecalhos(valores)
    mapeamento: dict[str, int] = {}
    for canonica, aliases in COLUNAS_CALIBRACAO.items():
        for alias in aliases:
            indice = indices.get(_normalizar_cabecalho(alias))
            if indice is not None:
                mapeamento[canonica] = indice
                break
    return mapeamento


def validar_estrutura_cal(colunas: Iterable[Any] | pd.DataFrame) -> bool:
    if isinstance(colunas, pd.DataFrame):
        valores = colunas.columns
    else:
        valores = colunas
    mapeamento = _mapear_colunas(valores)
    ausentes = [
        ROTULOS_COLUNAS[coluna]
        for coluna in COLUNAS_OBRIGATORIAS_CAL
        if coluna not in mapeamento
    ]
    if ausentes:
        logger.error(
            "Estrutura da aba CAL inválida; cabeçalhos obrigatórios ausentes: %s",
            ", ".join(ausentes),
        )
        return False
    return True


def _valor_celula(celula: Any) -> Any:
    valor = celula.value
    if celula.data_type == "f" or (isinstance(valor, str) and valor.startswith("=")):
        return None
    return valor


def _encontrar_linha_cabecalho(worksheet: Any) -> tuple[int, dict[str, int]]:
    for numero_linha, linha in enumerate(worksheet.iter_rows(), start=1):
        mapeamento = _mapear_colunas(_valor_celula(celula) for celula in linha)
        if all(coluna in mapeamento for coluna in COLUNAS_OBRIGATORIAS_CAL):
            return numero_linha, mapeamento
    logger.error("Estrutura da aba CAL inválida; cabeçalhos obrigatórios não encontrados.")
    raise ErroEstruturaCalibracao(
        "A estrutura da aba CAL foi alterada. Os cabeçalhos obrigatórios não foram encontrados."
    )


def _data_calibracao(valor: Any, epoch: datetime) -> date | None:
    if valor is None or valor is pd.NaT:
        return None
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    if isinstance(valor, (int, float)) and not isinstance(valor, bool):
        try:
            convertido = from_excel(valor, epoch)
        except (TypeError, ValueError, OverflowError):
            return None
        return convertido.date() if isinstance(convertido, datetime) else convertido
    if isinstance(valor, str):
        texto = valor.strip()
        if not texto:
            return None
        convertido = pd.to_datetime(texto, dayfirst=True, errors="coerce")
        if pd.isna(convertido):
            return None
        return convertido.date()
    return None


def _numero_dias(valor: Any) -> int | None:
    if valor is None or isinstance(valor, bool):
        return None
    try:
        numero = float(str(valor).strip().replace(",", "."))
    except (TypeError, ValueError):
        return None
    if not numero.is_integer() or numero <= 0:
        return None
    return int(numero)


def _mapa_periodicidade(workbook: Any) -> dict[str, int]:
    if "DADOS APOIO" not in workbook.sheetnames:
        return {}
    worksheet = workbook["DADOS APOIO"]
    cabecalho_encontrado: tuple[int, dict[str, int]] | None = None
    for numero_linha, linha in enumerate(worksheet.iter_rows(), start=1):
        indices = _indice_cabecalhos(_valor_celula(celula) for celula in linha)
        indice_equipamento = indices.get("EQUIPAMENTO")
        indice_validade = indices.get("VALIDADE CALIBRACAO")
        if indice_equipamento is not None and indice_validade is not None:
            cabecalho_encontrado = (
                numero_linha,
                {"equipamento": indice_equipamento, "validade": indice_validade},
            )
            break
    if cabecalho_encontrado is None:
        return {}

    linha_cabecalho, indices = cabecalho_encontrado
    periodicidades: dict[str, int] = {}
    for linha in worksheet.iter_rows(min_row=linha_cabecalho + 1):
        equipamento = _valor_celula(linha[indices["equipamento"]])
        dias = _numero_dias(_valor_celula(linha[indices["validade"]]))
        chave = normalizar_codigo(equipamento)
        if chave and dias is not None:
            periodicidades[chave] = dias
    return periodicidades


def calcular_dias_restantes(data_calibracao: Any, hoje: date | None = None) -> int | None:
    data_alvo = _data_calibracao(data_calibracao, datetime(1899, 12, 30))
    if data_alvo is None:
        return None
    return (data_alvo - (hoje or date.today())).days


def classificar_vencimento(
    proxima_calibracao: Any,
    hoje: date | None = None,
) -> str:
    dias_restantes = calcular_dias_restantes(proxima_calibracao, hoje)
    if dias_restantes is None:
        return "SEM PRÓXIMA CALIBRAÇÃO"
    if dias_restantes < 0:
        return "VENCIDO"
    if dias_restantes <= 30:
        return "A VENCER"
    return "NO PRAZO"


def _calcular_proxima_data(
    valor_proxima: Any,
    valor_data_calibracao: Any,
    periodicidades: dict[str, int],
    equipamento: Any,
    epoch: datetime,
) -> date | None:
    if valor_proxima is not None:
        return _data_calibracao(valor_proxima, epoch)

    data_base = _data_calibracao(valor_data_calibracao, epoch)
    dias = periodicidades.get(normalizar_codigo(equipamento))
    if data_base is None or dias is None:
        return None
    return data_base + timedelta(days=dias)


def _ler_aba_cal(workbook: Any) -> pd.DataFrame:
    nome_aba = next(
        (nome for nome in workbook.sheetnames if _normalizar_cabecalho(nome) == "CAL"),
        None,
    )
    if nome_aba is None:
        logger.error("Planilha de calibração sem a aba CAL.")
        raise ErroEstruturaCalibracao(
            'A aba "CAL" não foi encontrada no arquivo selecionado.'
        )
    worksheet = workbook[nome_aba]
    try:
        linha_cabecalho, indices = _encontrar_linha_cabecalho(worksheet)
    except ErroEstruturaCalibracao:
        validar_estrutura_cal(
            valor
            for linha in worksheet.iter_rows()
            for valor in (_valor_celula(celula) for celula in linha)
        )
        raise

    periodicidades = _mapa_periodicidade(workbook)
    indices_canonicos = {
        coluna: indices.get(coluna) for coluna in COLUNAS_CALIBRACAO
    }
    linhas: list[dict[str, Any]] = []
    for linha in worksheet.iter_rows(min_row=linha_cabecalho + 1):
        valores = {
            coluna: _valor_celula(linha[indice]) if indice is not None else None
            for coluna, indice in indices_canonicos.items()
        }
        codigo = valores["NUMERO_EQUIPAMENTO"]
        descricao = valores["DESCRICAO"]
        if all(valor is None or str(valor).strip() == "" for valor in (codigo, descricao)):
            continue
        data_prevista = _calcular_proxima_data(
            valores["PROXIMA_CALIBRACAO"],
            valores["DATA_CALIBRACAO"],
            periodicidades,
            descricao,
            workbook.epoch,
        )
        fabricante = valores["FABRICANTE"]
        if fabricante is None or str(fabricante).strip() in {"", "-"}:
            fabricante = "Não Definido"
        valores["FABRICANTE"] = fabricante
        valores["PROXIMA_CALIBRACAO"] = data_prevista
        valores["DIAS_RESTANTES"] = (
            (data_prevista - date.today()).days if data_prevista is not None else None
        )
        valores["STATUS_CALIBRACAO"] = classificar_vencimento(data_prevista)
        linhas.append(valores)

    dataframe = pd.DataFrame(linhas)
    for coluna in COLUNAS_CALIBRACAO:
        if coluna not in dataframe:
            dataframe[coluna] = pd.Series(dtype="object")
    if "DIAS_RESTANTES" not in dataframe:
        dataframe["DIAS_RESTANTES"] = pd.Series(dtype="Int64")

    dataframe = dataframe.rename(columns=ROTULOS_COLUNAS)
    dataframe["Código Normalizado"] = dataframe["Nº Equipamento"].apply(normalizar_codigo)
    dataframe["Próxima Calibração"] = pd.to_datetime(
        dataframe["Próxima Calibração"],
        errors="coerce",
    )
    dataframe["Dias Restantes"] = pd.to_numeric(
        dataframe["Dias Restantes"],
        errors="coerce",
    ).astype("Int64")
    for coluna in dataframe.columns:
        if coluna not in {"Próxima Calibração", "Dias Restantes"}:
            dataframe[coluna] = dataframe[coluna].where(dataframe[coluna].notna(), "")
    return dataframe


@st.cache_data(show_spinner=False)
def _carregar_calibracao_cache(conteudo: bytes) -> pd.DataFrame:
    with warnings.catch_warnings(record=True) as avisos:
        warnings.simplefilter("always", UserWarning)
        workbook = openpyxl.load_workbook(
            BytesIO(conteudo),
            data_only=False,
            read_only=False,
        )
    for aviso in avisos:
        logger.warning("Aviso ao abrir planilha de calibração: %s", aviso.message)
    try:
        return _ler_aba_cal(workbook)
    finally:
        workbook.close()


def carregar_calibracao(arquivo: Any) -> pd.DataFrame:
    try:
        if isinstance(arquivo, (bytes, bytearray)):
            conteudo = bytes(arquivo)
        elif hasattr(arquivo, "getvalue"):
            conteudo = arquivo.getvalue()
        elif hasattr(arquivo, "read"):
            conteudo = arquivo.read()
        else:
            raise TypeError("O arquivo de calibração não é um upload Excel válido.")
        if not conteudo:
            raise ValueError("O arquivo de calibração está vazio.")
        return _carregar_calibracao_cache(conteudo)
    except ErroEstruturaCalibracao:
        raise
    except (OSError, ValueError, TypeError, KeyError, openpyxl.utils.exceptions.InvalidFileException) as erro:
        logger.exception("Falha ao carregar a aba CAL.")
        raise ValueError(
            f"Não foi possível carregar a aba CAL: {erro}"
        ) from erro


def importar_planilha_cal(arquivo: Any, usuario: str) -> dict[str, int]:
    dataframe = carregar_calibracao(arquivo)
    return importar_equipamentos(dataframe, usuario)


def _valor_legivel(valor: Any) -> str:
    if valor is None or (not isinstance(valor, str) and pd.isna(valor)):
        return ""
    if isinstance(valor, (pd.Timestamp, datetime, date)):
        return pd.Timestamp(valor).strftime("%d/%m/%Y")
    return str(valor)


def _equipamentos_ativos(dataframe: pd.DataFrame) -> pd.Series:
    status = dataframe["Status Equipamento"].astype(str).map(_normalizar_cabecalho)
    return ~status.isin(STATUS_INATIVOS)


def _top_contagens(dataframe: pd.DataFrame, coluna: str, top: int = 10) -> pd.DataFrame:
    valores = (
        dataframe[coluna]
        .fillna("")
        .astype(str)
        .str.strip()
        .replace("", "Não Definido")
    )
    contagens = valores.value_counts().head(top).rename_axis(coluna).reset_index(name="Quantidade")
    return contagens.sort_values("Quantidade", ascending=True)


def _figura_barras_categoricas(dataframe: pd.DataFrame, coluna: str) -> go.Figure:
    figura = px.bar(
        dataframe,
        x="Quantidade",
        y=coluna,
        color=coluna,
        orientation="h",
        color_discrete_sequence=PALETA_GRAFICOS_CALIBRACAO,
    )
    figura.update_layout(
        showlegend=False,
        paper_bgcolor="rgba(0,0,0,0)",
        plot_bgcolor="rgba(0,0,0,0)",
        font_color="#f7f9fc",
        margin=dict(l=20, r=20, t=20, b=20),
    )
    return figura


def _figura_vencimentos_mensais(dataframe: pd.DataFrame) -> go.Figure:
    vencimentos = dataframe.loc[
        dataframe["Próxima Calibração"].notna(),
        ["Próxima Calibração", "Status Calibração"],
    ].copy()
    vencimentos["Mês"] = (
        vencimentos["Próxima Calibração"].dt.to_period("M").dt.to_timestamp()
    )
    mensal = (
        vencimentos.groupby(["Mês", "Status Calibração"])
        .size()
        .rename("Equipamentos")
        .reset_index()
        .sort_values("Mês")
    )
    figura = px.line(
        mensal,
        x="Mês",
        y="Equipamentos",
        color="Status Calibração",
        markers=True,
        category_orders={"Status Calibração": list(STATUS_CALIBRACAO)},
        color_discrete_map=CORES_STATUS_CALIBRACAO,
    )
    figura.update_layout(
        paper_bgcolor="rgba(0,0,0,0)",
        plot_bgcolor="rgba(0,0,0,0)",
        font_color="#f7f9fc",
        hovermode="x unified",
        xaxis_title="Mês da próxima calibração",
        yaxis_title="Quantidade de equipamentos",
        legend_title_text="Status",
    )
    return figura


def _exibir_kpi(coluna: Any, titulo: str, valor: Any, ajuda: str | None = None) -> None:
    coluna.metric(titulo, valor, help=ajuda)


def gerar_dashboard_calibracao(dataframe: pd.DataFrame) -> None:
    if dataframe.empty:
        st.info("A aba CAL não contém equipamentos para apresentar.")
        return

    ativo = _equipamentos_ativos(dataframe)
    total_ativos = int(ativo.sum())
    retirados = int(
        dataframe["Status Equipamento"]
        .astype(str)
        .map(_normalizar_cabecalho)
        .eq("RETIRADO DE USO")
        .sum()
    )
    vencidos = dataframe["Status Calibração"].eq("VENCIDO")
    a_vencer = dataframe["Status Calibração"].eq("A VENCER")
    no_prazo = dataframe["Status Calibração"].eq("NO PRAZO")
    equipamentos_ok = int((ativo & no_prazo).sum())
    conformidade = equipamentos_ok / total_ativos * 100 if total_ativos else 0.0
    fabricantes_indefinidos = int(
        dataframe["Fabricante"].astype(str).eq("Não Definido").sum()
    )
    sem_data = int(dataframe["Próxima Calibração"].isna().sum())
    duplicados_mask = (
        dataframe["Código Normalizado"].ne("")
        & dataframe["Código Normalizado"].duplicated(keep=False)
    )
    duplicados = int(duplicados_mask.sum())

    kpis = st.columns(6)
    _exibir_kpi(kpis[0], "Total Equipamentos", len(dataframe))
    _exibir_kpi(kpis[1], "No Prazo", int(no_prazo.sum()))
    _exibir_kpi(kpis[2], "A Vencer (30 dias)", int(a_vencer.sum()))
    _exibir_kpi(kpis[3], "Vencidos", int(vencidos.sum()))
    _exibir_kpi(kpis[4], "Retirados de Uso", retirados)
    _exibir_kpi(kpis[5], "Conformidade", f"{conformidade:.1f}%")

    extras = st.columns(3)
    _exibir_kpi(extras[0], "Fabricantes Não Definidos", fabricantes_indefinidos)
    _exibir_kpi(extras[1], "Sem Próxima Calibração", sem_data)
    _exibir_kpi(extras[2], "Equipamentos Duplicados", duplicados)

    setor, fabricante = st.columns(2)
    por_setor = _top_contagens(dataframe, "Setor", 20)
    por_fabricante = _top_contagens(dataframe, "Fabricante", 20)
    setor.write("Equipamentos por Setor")
    setor.plotly_chart(
        _figura_barras_categoricas(por_setor, "Setor"),
        use_container_width=True,
    )
    fabricante.write("Equipamentos por Fabricante")
    fabricante.plotly_chart(
        _figura_barras_categoricas(por_fabricante, "Fabricante"),
        use_container_width=True,
    )

    vencimentos = dataframe.loc[dataframe["Próxima Calibração"].notna()]
    if vencimentos.empty:
        st.info("Não há datas de calibração disponíveis para o gráfico mensal.")
    else:
        st.subheader("Vencimentos por Mês")
        st.plotly_chart(
            _figura_vencimentos_mensais(dataframe),
            use_container_width=True,
        )

    extra_esquerda, extra_direita = st.columns(2)
    extra_esquerda.subheader("Top 10 Equipamentos por Tipo")
    extra_esquerda.plotly_chart(
        _figura_barras_categoricas(_top_contagens(dataframe, "Descrição"), "Descrição"),
        use_container_width=True,
    )
    extra_direita.subheader("Top 10 Setores com mais Equipamentos")
    extra_direita.plotly_chart(
        _figura_barras_categoricas(_top_contagens(dataframe, "Setor"), "Setor"),
        use_container_width=True,
    )


def _filtrar_equipamentos(dataframe: pd.DataFrame, prefixo: str) -> pd.DataFrame:
    filtros = st.columns(4)
    filtro_codigo = filtros[0].text_input("Número Equipamento", key=f"{prefixo}_numero")
    filtro_descricao = filtros[1].text_input("Descrição", key=f"{prefixo}_descricao")
    setores = ["Todos", *sorted(dataframe["Setor"].astype(str).replace("", "Não Definido").unique())]
    fabricantes = [
        "Todos",
        *sorted(dataframe["Fabricante"].astype(str).replace("", "Não Definido").unique()),
    ]
    filtro_setor = filtros[2].selectbox("Setor", setores, key=f"{prefixo}_setor")
    filtro_fabricante = filtros[3].selectbox(
        "Fabricante",
        fabricantes,
        key=f"{prefixo}_fabricante",
    )

    segunda_linha = st.columns(3)
    filtro_status_cal = segunda_linha[0].selectbox(
        "Status Calibração",
        ["Todos", *STATUS_CALIBRACAO],
        key=f"{prefixo}_status_cal",
    )
    status_equipamentos = sorted(
        dataframe["Status Equipamento"].astype(str).replace("", "Não Definido").unique()
    )
    filtro_status_equip = segunda_linha[1].selectbox(
        "Status Equipamento",
        ["Todos", *status_equipamentos],
        key=f"{prefixo}_status_equip",
    )
    locais = sorted(dataframe["Local de Uso"].astype(str).replace("", "Não Definido").unique())
    filtro_local = segunda_linha[2].selectbox(
        "Local de Uso",
        ["Todos", *locais],
        key=f"{prefixo}_local",
    )

    filtrado = dataframe.copy()
    codigo_digitado = normalizar_codigo(filtro_codigo)
    if codigo_digitado:
        filtrado = filtrado[
            filtrado["Código Normalizado"].str.contains(
                re.escape(codigo_digitado),
                na=False,
            )
        ]
    if filtro_descricao.strip():
        filtrado = filtrado[
            filtrado["Descrição"].astype(str).str.contains(
                re.escape(filtro_descricao.strip()),
                case=False,
                na=False,
            )
        ]
    for coluna, filtro in (
        ("Setor", filtro_setor),
        ("Fabricante", filtro_fabricante),
        ("Status Calibração", filtro_status_cal),
        ("Status Equipamento", filtro_status_equip),
        ("Local de Uso", filtro_local),
    ):
        if filtro != "Todos":
            filtrado = filtrado[filtrado[coluna].astype(str).replace("", "Não Definido") == filtro]
    return filtrado


def exportar_excel_calibracao(dataframe: pd.DataFrame, nome_aba: str = "Calibracao") -> bytes:
    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as escritor:
        dataframe.drop(columns=["Código Normalizado"], errors="ignore").to_excel(
            escritor,
            index=False,
            sheet_name=nome_aba[:31],
        )
        worksheet = escritor.sheets[nome_aba[:31]]
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for coluna in worksheet.columns:
            letra = coluna[0].column_letter
            largura = max(
                (len(str(celula.value or "")) for celula in coluna),
                default=10,
            )
            worksheet.column_dimensions[letra].width = min(max(largura + 2, 12), 48)
    return buffer.getvalue()


def _exportar_csv_calibracao(dataframe: pd.DataFrame) -> bytes:
    return dataframe.drop(
        columns=["Código Normalizado"],
        errors="ignore",
    ).to_csv(index=False).encode("utf-8-sig")


def _texto_pdf(texto: Any) -> str:
    dados = _valor_legivel(texto).encode("cp1252", errors="replace")
    resultado = []
    for byte in dados:
        if byte in (ord("("), ord(")"), ord("\\")):
            resultado.append("\\" + chr(byte))
        elif 32 <= byte <= 126:
            resultado.append(chr(byte))
        else:
            resultado.append(f"\\{byte:03o}")
    return "".join(resultado)


def _pdf_calibracao(dataframe: pd.DataFrame) -> bytes:
    writer = PdfWriter()
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
            NameObject("/Encoding"): NameObject("/WinAnsiEncoding"),
        }
    )
    font_ref = writer._add_object(font)
    largura_pagina, altura_pagina = 842, 595
    colunas = [coluna for coluna in CAMPOS_RELATORIO_CAL if coluna in dataframe.columns]
    linhas: list[list[str]] = [
        [_valor_legivel(valor) for valor in linha]
        for linha in dataframe[colunas].itertuples(index=False, name=None)
    ]
    larguras = [13, 24, 9, 15, 18, 13, 14, 14, 18, 20]
    larguras = larguras[: len(colunas)]
    limite_linhas = 29
    paginas = [
        linhas[inicio : inicio + limite_linhas]
        for inicio in range(0, len(linhas), limite_linhas)
    ] or [[]]

    for numero_pagina, linhas_pagina in enumerate(paginas, start=1):
        comandos = [
            "BT /F1 15 Tf 30 560 Td (Relatorio de Gestao de Calibracao) Tj ET",
            f"BT /F1 8 Tf 30 542 Td (Pagina {numero_pagina} de {len(paginas)} | Registros: {len(linhas)}) Tj ET",
        ]
        y = 520
        comandos.append(f"BT /F1 7 Tf 30 {y} Td ({_texto_pdf(' | '.join(colunas))}) Tj ET")
        y -= 15
        for linha in linhas_pagina:
            celulas = []
            for valor, largura in zip(linha, larguras):
                celulas.append(valor.replace("\r", " ").replace("\n", " ")[:largura])
            texto = " | ".join(celulas)
            comandos.append(f"BT /F1 6 Tf 30 {y} Td ({_texto_pdf(texto)}) Tj ET")
            y -= 15

        pagina = writer.add_blank_page(width=largura_pagina, height=altura_pagina)
        pagina[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font_ref})}
        )
        conteudo = DecodedStreamObject()
        conteudo.set_data(("\n".join(comandos)).encode("ascii"))
        pagina[NameObject("/Contents")] = writer._add_object(conteudo)

    buffer = BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def gerar_alertas_calibracao(dataframe: pd.DataFrame) -> None:
    categorias = (
        ("🔴 Vencidos", dataframe["Dias Restantes"].lt(0)),
        (
            "🟡 Próximos 30 dias",
            dataframe["Dias Restantes"].between(0, 30, inclusive="both"),
        ),
        (
            "🔵 Próximos 60 dias",
            dataframe["Dias Restantes"].between(31, 60, inclusive="both"),
        ),
    )
    for titulo, mascara in categorias:
        grupo = dataframe.loc[mascara].sort_values(
            "Próxima Calibração",
            na_position="last",
        )
        st.subheader(f"{titulo} — {len(grupo)} equipamento(s)")
        if grupo.empty:
            st.success("Nenhum equipamento nesta categoria.")
            continue
        st.dataframe(
            grupo.drop(columns=["Código Normalizado"], errors="ignore"),
            hide_index=True,
            use_container_width=True,
        )
        st.download_button(
            f"Exportar {titulo}",
            exportar_excel_calibracao(grupo, "Alertas"),
            file_name=f"alertas_{titulo.split()[-1].lower()}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            key=f"exportar_alerta_{titulo}",
        )


def gerar_relatorio_calibracao(dataframe: pd.DataFrame) -> None:
    st.subheader("Relatórios de calibração")
    criterio = st.selectbox(
        "Filtro do relatório",
        ["Todos", "Vencidos", "A Vencer", "No Prazo", "Por Setor", "Por Fabricante"],
        key="filtro_relatorio_cal",
    )
    filtrado = dataframe.copy()
    if criterio in {"Vencidos", "A Vencer", "No Prazo"}:
        filtrado = _filtrar_status_relatorio(filtrado, criterio)
    elif criterio == "Por Setor":
        setor = st.selectbox(
            "Setor",
            sorted(filtrado["Setor"].astype(str).replace("", "Não Definido").unique()),
            key="relatorio_cal_setor",
        )
        filtrado = filtrado[
            filtrado["Setor"].astype(str).replace("", "Não Definido") == setor
        ]
    elif criterio == "Por Fabricante":
        fabricante = st.selectbox(
            "Fabricante",
            sorted(filtrado["Fabricante"].astype(str).replace("", "Não Definido").unique()),
            key="relatorio_cal_fabricante",
        )
        filtrado = filtrado[
            filtrado["Fabricante"].astype(str).replace("", "Não Definido")
            == fabricante
        ]

    relatorio = filtrado.reindex(columns=CAMPOS_RELATORIO_CAL)
    if st.button("GERAR REPORT CAL", type="primary", key="gerar_report_cal"):
        st.session_state["_relatorio_cal_gerado"] = relatorio.copy()

    if "_relatorio_cal_gerado" not in st.session_state:
        st.caption("Selecione um filtro e gere o relatório para habilitar as exportações.")
        return

    resultado = st.session_state["_relatorio_cal_gerado"]
    st.caption(f"Relatório pronto: {len(resultado)} registro(s).")
    st.dataframe(resultado, hide_index=True, use_container_width=True)
    downloads = st.columns(3)
    downloads[0].download_button(
        "Excel",
        exportar_excel_calibracao(resultado, "Report CAL"),
        "report_cal.xlsx",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        key="download_report_cal_excel",
    )
    downloads[1].download_button(
        "CSV",
        _exportar_csv_calibracao(resultado),
        "report_cal.csv",
        "text/csv",
        key="download_report_cal_csv",
    )
    downloads[2].download_button(
        "PDF",
        _pdf_calibracao(resultado),
        "report_cal.pdf",
        "application/pdf",
        key="download_report_cal_pdf",
    )


def _filtrar_status_relatorio(
    dataframe: pd.DataFrame,
    criterio: str,
) -> pd.DataFrame:
    status = {
        "Vencidos": "VENCIDO",
        "A Vencer": "A VENCER",
        "No Prazo": "NO PRAZO",
    }.get(criterio)
    if status is None:
        raise ValueError(f"Filtro de relatório desconhecido: {criterio}.")
    return dataframe[dataframe["Status Calibração"] == status]


def _gerar_tela_equipamentos(dataframe: pd.DataFrame) -> None:
    filtrado = _filtrar_equipamentos(dataframe, "equipamentos_cal")
    st.caption(f"{len(filtrado)} equipamento(s) encontrado(s).")
    st.dataframe(
        filtrado.drop(columns=["Código Normalizado"], errors="ignore"),
        hide_index=True,
        use_container_width=True,
    )
    st.download_button(
        "Exportar equipamentos para Excel",
        exportar_excel_calibracao(filtrado, "Equipamentos"),
        "equipamentos_calibracao.xlsx",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        key="exportar_equipamentos_cal",
    )
    if dataframe.empty:
        return
    equipamentos_disponiveis = dataframe[
        ["Nº Equipamento", "Descrição", "Código Normalizado"]
    ].copy()
    equipamentos_disponiveis["Opção"] = (
        equipamentos_disponiveis["Nº Equipamento"].astype(str)
        + " — "
        + equipamentos_disponiveis["Descrição"].astype(str)
    )
    opcoes = dict(
        zip(
            equipamentos_disponiveis["Opção"],
            equipamentos_disponiveis["Código Normalizado"],
        )
    )
    escolha = st.selectbox(
        "Abrir equipamento para ver detalhes e histórico",
        list(opcoes),
        key="equipamento_cal_historico",
    )
    selecionado = dataframe[
        dataframe["Código Normalizado"] == opcoes[escolha]
    ].drop(columns=["Código Normalizado"], errors="ignore")
    aba_dados, aba_historico = st.tabs(["Equipamento", "📜 Histórico"])
    with aba_dados:
        st.dataframe(
            selecionado.T.rename(columns={selecionado.index[0]: "Valor"}),
            use_container_width=True,
        )
    with aba_historico:
        historico = carregar_historico_equipamento(opcoes[escolha])
        if historico.empty:
            st.info("Ainda não há alterações registradas para este equipamento.")
        else:
            st.dataframe(
                _formatar_data_hora_local(historico, "Data"),
                hide_index=True,
                use_container_width=True,
            )


def _filtrar_por_periodo(
    dataframe: pd.DataFrame,
    inicio: date,
    fim: date,
) -> pd.DataFrame:
    if dataframe.empty or "Data/Hora" not in dataframe:
        return dataframe.copy()
    datas = (
        pd.to_datetime(dataframe["Data/Hora"], errors="coerce", utc=True)
        .dt.tz_convert("America/Sao_Paulo")
        .dt.date
    )
    return dataframe.loc[(datas >= inicio) & (datas <= fim)].copy()


def _formatar_data_hora_local(dataframe: pd.DataFrame, coluna: str) -> pd.DataFrame:
    resultado = dataframe.copy()
    if coluna in resultado and not resultado.empty:
        resultado[coluna] = (
            pd.to_datetime(resultado[coluna], errors="coerce", utc=True)
            .dt.tz_convert("America/Sao_Paulo")
            .dt.strftime("%d/%m/%Y %H:%M:%S")
        )
    return resultado


def _filtrar_valor(
    dataframe: pd.DataFrame,
    coluna: str,
    valor: str,
) -> pd.DataFrame:
    if valor == "Todos" or coluna not in dataframe:
        return dataframe
    return dataframe[dataframe[coluna].astype(str) == valor]


def _gerar_tela_auditoria(usuario: dict[str, Any]) -> None:
    if usuario["perfil"] not in PAPEIS_AUDITORIA:
        st.error("Seu perfil não tem permissão para acessar a auditoria.")
        return
    st.subheader("📜 Auditoria de alterações")
    try:
        alteracoes = carregar_auditoria_alteracoes()
        eventos = carregar_auditoria_eventos()
    except SQLAlchemyError as erro:
        logger.exception("Falha ao consultar auditoria de calibração.")
        st.error(f"Não foi possível carregar a auditoria: {erro}")
        return

    data_fim = date.today()
    data_inicio = data_fim - timedelta(days=365)
    periodo = st.date_input(
        "Período",
        value=(data_inicio, data_fim),
        key="periodo_auditoria_cal",
    )
    if not isinstance(periodo, tuple) or len(periodo) != 2:
        st.info("Selecione a data inicial e final para filtrar a auditoria.")
        return
    inicio, fim = periodo
    if inicio > fim:
        st.warning("A data inicial deve ser anterior à data final.")
        return

    aba_alteracoes, aba_eventos = st.tabs(
        ["Alterações por campo", "Eventos do sistema"]
    )
    with aba_alteracoes:
        filtrado = _filtrar_por_periodo(alteracoes, inicio, fim)
        filtros = st.columns(3)
        for coluna, rotulo in (
            ("Usuário", "Usuário"),
            ("Equipamento", "Equipamento"),
            ("Campo Alterado", "Campo"),
            ("Tipo de Alteração", "Tipo de Alteração"),
        ):
            valores = (
                ["Todos", *sorted(filtrado[coluna].dropna().astype(str).unique())]
                if coluna in filtrado
                else ["Todos"]
            )
            indice = (
                ("Usuário", "Equipamento", "Campo Alterado", "Tipo de Alteração")
                .index(coluna)
            )
            escolha = filtros[indice % len(filtros)].selectbox(
                rotulo,
                valores,
                key=f"filtro_auditoria_cal_{coluna}",
            )
            filtrado = _filtrar_valor(filtrado, coluna, escolha)
        alteracoes_exibicao = _formatar_data_hora_local(filtrado, "Data/Hora")
        st.caption(f"{len(filtrado)} alteração(ões) encontrada(s).")
        st.dataframe(alteracoes_exibicao, hide_index=True, use_container_width=True)
        exportacoes = st.columns(2)
        exportacoes[0].download_button(
            "Exportar alterações para Excel",
            exportar_dataframe_excel(alteracoes_exibicao, "Auditoria"),
            "auditoria_calibracao.xlsx",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            key="exportar_auditoria_cal_excel",
        )
        exportacoes[1].download_button(
            "Exportar alterações para CSV",
            exportar_dataframe_csv(alteracoes_exibicao),
            "auditoria_calibracao.csv",
            "text/csv",
            key="exportar_auditoria_cal_csv",
        )

    with aba_eventos:
        filtrado_eventos = eventos.copy()
        if not filtrado_eventos.empty:
            datas_eventos = pd.to_datetime(
                filtrado_eventos["Data/Hora"],
                errors="coerce",
                utc=True,
            ).dt.tz_convert("America/Sao_Paulo").dt.date
            filtrado_eventos = filtrado_eventos.loc[
                (datas_eventos >= inicio) & (datas_eventos <= fim)
            ]
        filtros_eventos = st.columns(3)
        for coluna, rotulo in (
            ("Usuário", "Usuário"),
            ("Equipamento", "Equipamento"),
            ("Tipo de Evento", "Tipo de Alteração"),
        ):
            valores = (
                ["Todos", *sorted(filtrado_eventos[coluna].dropna().astype(str).unique())]
                if coluna in filtrado_eventos
                else ["Todos"]
            )
            escolha = filtros_eventos[
                ("Usuário", "Equipamento", "Tipo de Evento").index(coluna)
            ].selectbox(
                rotulo,
                valores,
                key=f"filtro_eventos_cal_{coluna}",
            )
            filtrado_eventos = _filtrar_valor(
                filtrado_eventos,
                coluna,
                escolha,
            )
        eventos_exibicao = _formatar_data_hora_local(
            filtrado_eventos,
            "Data/Hora",
        )
        st.caption(f"{len(filtrado_eventos)} evento(s) encontrado(s).")
        st.dataframe(eventos_exibicao, hide_index=True, use_container_width=True)
        exportacoes = st.columns(2)
        exportacoes[0].download_button(
            "Exportar eventos para Excel",
            exportar_dataframe_excel(eventos_exibicao, "Eventos"),
            "eventos_auditoria_calibracao.xlsx",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            key="exportar_eventos_cal_excel",
        )
        exportacoes[1].download_button(
            "Exportar eventos para CSV",
            exportar_dataframe_csv(eventos_exibicao),
            "eventos_auditoria_calibracao.csv",
            "text/csv",
            key="exportar_eventos_cal_csv",
        )


def _gerar_tela_solicitacoes_acesso(usuario: dict[str, Any]) -> None:
    if usuario["perfil"] != "ADMIN":
        st.error("Somente o perfil ADMIN pode administrar acessos.")
        return
    st.subheader("Solicitações de acesso")
    try:
        pendentes = listar_usuarios("PENDENTE")
        if pendentes.empty:
            st.info("Não há solicitações pendentes.")
        else:
            st.dataframe(
                pendentes[["nome", "email", "criado_em"]],
                hide_index=True,
                use_container_width=True,
            )
            rotulos = {
                f"{linha.id} — {linha.nome} ({linha.email})": int(linha.id)
                for linha in pendentes.itertuples(index=False)
            }
            escolha = st.selectbox(
                "Solicitação",
                list(rotulos),
                key="solicitacao_cal_pendente",
            )
            perfil = st.selectbox(
                "Perfil a conceder",
                ["QUALIDADE", "LEITOR", "ADMIN"],
                key="perfil_solicitacao_cal",
            )
            aprovar, rejeitar = st.columns(2)
            if aprovar.button(
                "Aprovar solicitação",
                type="primary",
                key="aprovar_usuario_cal",
                use_container_width=True,
            ):
                aprovar_solicitacao(
                    rotulos[escolha],
                    usuario["email"],
                    perfil,
                )
                st.success("Usuário aprovado.")
                st.rerun()
            if rejeitar.button(
                "Rejeitar solicitação",
                key="rejeitar_usuario_cal",
                use_container_width=True,
            ):
                rejeitar_solicitacao(rotulos[escolha], usuario["email"])
                st.success("Solicitação rejeitada.")
                st.rerun()
    except (SQLAlchemyError, ValueError, PermissionError) as erro:
        logger.exception("Falha ao administrar solicitações de acesso.")
        st.error(f"Não foi possível concluir a operação de solicitações: {erro}")


def _gerar_tela_usuarios(usuario: dict[str, Any]) -> None:
    if usuario["perfil"] != "ADMIN":
        st.error("Somente o perfil ADMIN pode administrar acessos.")
        return
    st.subheader("Usuários")
    try:
        ativos = listar_usuarios("ATIVO")
        if ativos.empty:
            st.info("Não há usuários ativos.")
            return
        st.dataframe(
            ativos[["nome", "email", "perfil", "aprovado_em"]],
            hide_index=True,
            use_container_width=True,
        )
    except (SQLAlchemyError, ValueError, PermissionError) as erro:
        logger.exception("Falha ao consultar usuários ativos.")
        st.error(f"Não foi possível carregar os usuários: {erro}")


def _gerar_tela_perfis_permissoes(usuario: dict[str, Any]) -> None:
    if usuario["perfil"] != "ADMIN":
        st.error("Somente o perfil ADMIN pode administrar permissões.")
        return
    st.subheader("Perfis e permissões")
    try:
        ativos = listar_usuarios("ATIVO")
        if ativos.empty:
            st.info("Não há usuários ativos.")
            return
        alteraveis = ativos[ativos["email"] != usuario["email"]]
        if alteraveis.empty:
            st.caption("Não há outros usuários ativos para alterar.")
            return
        opcoes = {
            f"{linha.id} — {linha.nome} ({linha.email})": linha
            for linha in alteraveis.itertuples(index=False)
        }
        selecionado = st.selectbox(
            "Usuário",
            list(opcoes),
            key="usuario_cal_alterar_perfil",
        )
        novo_perfil = st.selectbox(
            "Perfil",
            list(PERFIS_USUARIO),
            index=list(PERFIS_USUARIO).index(opcoes[selecionado].perfil),
            key="novo_perfil_cal",
        )
        if st.button("Salvar perfil", key="salvar_perfil_cal"):
            alterar_perfil_usuario(
                int(opcoes[selecionado].id),
                usuario["email"],
                novo_perfil,
            )
            st.success("Perfil atualizado.")
            st.rerun()
    except (SQLAlchemyError, ValueError, PermissionError) as erro:
        logger.exception("Falha ao alterar perfil de usuário.")
        st.error(f"Não foi possível atualizar o perfil: {erro}")


def _gerar_tela_administracao_global(usuario: dict[str, Any]) -> None:
    if usuario["perfil"] != "ADMIN":
        st.error("Somente o perfil ADMIN pode acessar a Administração.")
        return
    st.markdown("## ⚙ Administração")
    usuarios, solicitacoes, permissoes, auditoria, notificacoes, logs, configuracoes = (
        st.tabs(
            [
                "Usuários",
                "Solicitações de Acesso",
                "Perfis e Permissões",
                "Auditoria Geral",
                "Notificações",
                "Logs",
                "Configurações",
            ]
        )
    )
    with usuarios:
        _gerar_tela_usuarios(usuario)
    with solicitacoes:
        _gerar_tela_solicitacoes_acesso(usuario)
    with permissoes:
        _gerar_tela_perfis_permissoes(usuario)
    with auditoria:
        st.subheader("Auditoria Geral")
        try:
            calibracao = carregar_auditoria_alteracoes()
            eventos = carregar_auditoria_eventos()
            from database.cn_store import carregar_auditoria_administracao

            cns = carregar_auditoria_administracao(usuario["email"])
        except (SQLAlchemyError, PermissionError) as erro:
            logger.exception("Falha ao carregar auditoria geral.")
            st.error(f"Não foi possível carregar a auditoria geral: {erro}")
        else:
            auditoria_geral = pd.concat(
                [
                    calibracao.assign(Módulo="Calibração"),
                    eventos.assign(Módulo="Eventos"),
                    cns.assign(Módulo="CN"),
                ],
                ignore_index=True,
                sort=False,
            )
            st.dataframe(
                auditoria_geral,
                hide_index=True,
                use_container_width=True,
            )
    with notificacoes:
        st.subheader("Notificações")
        try:
            from database.cn_store import carregar_todas_notificacoes_administracao

            notificacoes_cns = carregar_todas_notificacoes_administracao(
                usuario["email"]
            )
        except (SQLAlchemyError, PermissionError) as erro:
            logger.exception("Falha ao carregar notificações globais.")
            st.error(f"Não foi possível carregar as notificações: {erro}")
        else:
            if notificacoes_cns.empty:
                st.info("Não há notificações administrativas registradas.")
            else:
                st.dataframe(
                    notificacoes_cns,
                    hide_index=True,
                    use_container_width=True,
                )
    with logs:
        st.subheader("Logs")
        st.info("Não há um repositório de logs de aplicação configurado.")
    with configuracoes:
        st.subheader("Configurações")
        st.info("Não há configurações globais persistidas para administrar.")


def _gerar_tela_administracao_usuarios(usuario: dict[str, Any]) -> None:
    _gerar_tela_solicitacoes_acesso(usuario)
    _gerar_tela_usuarios(usuario)
    _gerar_tela_perfis_permissoes(usuario)


def _render_autenticacao() -> dict[str, Any] | None:
    return render_autenticacao()


def gerar_dashboard(dataframe: pd.DataFrame) -> None:
    gerar_dashboard_calibracao(dataframe)


def gerar_alertas(dataframe: pd.DataFrame) -> None:
    gerar_alertas_calibracao(dataframe)


def gerar_relatorios(dataframe: pd.DataFrame) -> None:
    gerar_relatorio_calibracao(dataframe)


def render_gestao_calibracao() -> None:
    st.markdown("## 📏 Gestão de Calibração")
    usuario = _render_autenticacao()
    if usuario is None:
        return
    titulo_usuario, botao_sair = st.columns([4, 1])
    titulo_usuario.caption(
        f"Sessão: {usuario['nome']} · Perfil {usuario['perfil']}"
    )
    if botao_sair.button(
        "Sair",
        key="sair_gestao_calibracao",
        use_container_width=True,
    ):
        st.session_state.pop("_usuario_gestao_calibracao", None)
        st.rerun()

    arquivo = None
    if usuario["perfil"] in PAPEIS_AUDITORIA:
        arquivo = st.file_uploader(
            "Importe a planilha Excel com a aba CAL",
            type=["xlsx", "xlsm"],
            key="arquivo_gestao_calibracao",
        )
    if arquivo is not None:
        conteudo = arquivo.getvalue()
        importacao_id = f"{arquivo.name}:{len(conteudo)}:{hashlib.sha256(conteudo).hexdigest()}"
        if st.session_state.get("_importacao_calibracao_id") != importacao_id:
            try:
                resultado_importacao = importar_planilha_cal(
                    conteudo,
                    usuario["email"],
                )
            except ErroEstruturaCalibracao as erro:
                logger.exception("Estrutura inválida ao importar calibração.")
                st.error("A estrutura da aba CAL foi alterada.")
                st.caption(str(erro))
            except (ValueError, OSError, TypeError, KeyError, SQLAlchemyError, PermissionError) as erro:
                logger.exception("Falha ao importar a planilha de calibração.")
                st.error(f"Não foi possível importar a planilha de calibração: {erro}")
            else:
                st.session_state["_importacao_calibracao_id"] = importacao_id
                st.session_state["_resultado_importacao_calibracao"] = resultado_importacao
        resultado_importacao = st.session_state.get(
            "_resultado_importacao_calibracao"
        )
        if resultado_importacao and st.session_state.get("_importacao_calibracao_id") == importacao_id:
            st.success(
                "Importação concluída: "
                f"{resultado_importacao['novos']} novo(s), "
                f"{resultado_importacao['atualizados']} atualizado(s), "
                f"{resultado_importacao['inalterados']} sem alteração."
            )
    try:
        dataframe = carregar_equipamentos()
    except SQLAlchemyError as erro:
        logger.exception("Falha ao consultar equipamentos de calibração.")
        st.error(f"Não foi possível carregar o histórico de calibração: {erro}")
        return
    if dataframe.empty:
        st.info("Ainda não há equipamentos cadastrados; importe uma planilha com a aba CAL.")

    opcoes_menu = ["Dashboard", "Equipamentos", "Alertas", "Relatórios"]
    if usuario["perfil"] in {"ADMIN", "QUALIDADE"}:
        opcoes_menu.append("📜 Auditoria")
    menu = st.radio(
        "Menu Gestão de Calibração",
        opcoes_menu,
        horizontal=True,
        key="menu_gestao_calibracao",
        label_visibility="collapsed",
    )
    if dataframe.empty and menu in {
        "Dashboard",
        "Equipamentos",
        "Alertas",
        "Relatórios",
    }:
        st.info("Este painel será exibido após a primeira importação válida.")
        return
    if menu == "Dashboard":
        gerar_dashboard(dataframe)
    elif menu == "Equipamentos":
        st.subheader("Equipamentos")
        _gerar_tela_equipamentos(dataframe)
    elif menu == "Alertas":
        gerar_alertas(dataframe)
    elif menu == "Relatórios":
        gerar_relatorios(dataframe)
    elif menu == "📜 Auditoria":
        _gerar_tela_auditoria(usuario)


def render_administracao_global() -> None:
    from core.administration import render_administracao_global as renderizar

    renderizar()


__all__ = [
    "ErroEstruturaCalibracao",
    "aprovar_solicitacao",
    "autenticar_usuario",
    "calcular_dias_restantes",
    "carregar_calibracao",
    "classificar_vencimento",
    "criar_usuario",
    "exportar_excel_calibracao",
    "gerar_alertas",
    "gerar_alertas_calibracao",
    "gerar_dashboard",
    "gerar_dashboard_calibracao",
    "gerar_relatorios",
    "gerar_relatorio_calibracao",
    "importar_planilha_cal",
    "inicializar_banco",
    "normalizar_codigo",
    "registrar_auditoria",
    "render_administracao_global",
    "render_gestao_calibracao",
    "validar_estrutura_cal",
]
