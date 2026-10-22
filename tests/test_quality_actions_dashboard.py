from datetime import date
from io import BytesIO
from types import SimpleNamespace

import openpyxl
import pandas as pd
import unittest

from core.quality_actions_dashboard import (
    calcular_distribuicao_target,
    calcular_indicadores_acoes,
    calcular_meta_esperada,
    calcular_situacao_prazo,
    ler_target_gerencia_excel,
)


class QualityActionsDashboardTests(unittest.TestCase):
    def test_calcular_situacao_prazo_recalcula_abertas_e_preserva_concluidas(self):
        dataframe = pd.DataFrame(
            {
                "Status": ["Aberto", "Aberto", "Aberto", "Concluído"],
                "Prazo": ["2026-10-19", "2026-10-20", None, "2026-10-01"],
            }
        )

        situacoes = calcular_situacao_prazo(
            dataframe, "Status", "Prazo", date(2026, 10, 20)
        )

        self.assertEqual(
            situacoes.tolist(),
            ["Vencida", "No prazo", "Sem prazo", "Concluída"],
        )

    def test_calcular_meta_esperada_aplica_proporcao_calendario(self):
        casos = [
            ("SETEMBRO", 560.0),
            ("OUTUBRO", 560 * 20 / 31),
            ("NOVEMBRO", 0.0),
        ]
        for mes, esperada in casos:
            with self.subTest(mes=mes):
                self.assertAlmostEqual(
                    calcular_meta_esperada(560, mes, date(2026, 10, 20)),
                    esperada,
                )

    def test_calcular_indicadores_acoes_apura_kpis_e_alertas(self):
        dataframe = pd.DataFrame(
            {"Status": ["Concluído", "Aberto", "Aberto", "Outro"]}
        )
        situacoes = pd.Series(["Concluída", "Vencida", "No prazo", "Sem prazo"])

        indicadores = calcular_indicadores_acoes(
            dataframe,
            "Status",
            situacoes,
            10,
            "OUTUBRO",
            date(2026, 10, 20),
        )

        self.assertEqual(indicadores["total_realizado"], 4)
        self.assertEqual(indicadores["concluidas"], 1)
        self.assertEqual(indicadores["em_aberto"], 2)
        self.assertEqual(indicadores["vencidas"], 1)
        self.assertEqual(indicadores["no_prazo"], 1)
        self.assertAlmostEqual(indicadores["percentual_atingimento"], 0.4)
        self.assertAlmostEqual(indicadores["falta_para_meta"], 6)
        self.assertAlmostEqual(indicadores["meta_esperada"], 10 * 20 / 31)
        self.assertAlmostEqual(
            indicadores["diferenca_meta"], 4 - 10 * 20 / 31
        )

    def test_distribui_meta_por_peso_e_quantidade_e_recalcula_realizado(self):
        linhas = pd.DataFrame(
            {
                "Responsável": ["A", "B"],
                "Peso": [1, 3],
                "Qtd. M.O": [10, 10],
                "Realizado WK 1": [2, 4],
            }
        )

        resultado = calcular_distribuicao_target(linhas, 100)

        self.assertEqual(resultado["Target Mensal Calculado"].tolist(), [25, 75])
        self.assertEqual(resultado["Total Realizado Calculado"].tolist(), [2, 4])
        self.assertEqual(resultado["Pendente Calculado"].tolist(), [23, 71])

    def test_pendente_calculado_arredonda_como_formula_excel(self):
        linhas = pd.DataFrame(
            {
                "Responsável": ["A", "B"],
                "Peso": [1, 1],
                "Qtd. M.O": [1, 1],
                "Realizado WK 1": [0, 0],
            }
        )

        resultado = calcular_distribuicao_target(linhas, 1)

        self.assertEqual(resultado["Pendente Calculado"].tolist(), [1, 1])

    def test_ler_target_gerencia_ignora_formulas_e_recalcula_blocos(self):
        workbook = openpyxl.Workbook()
        worksheet = workbook.active
        worksheet.title = "Target Gerencia Outubro"
        worksheet["Y2"] = 2026
        worksheet["Z2"] = 100
        header = [
            "Reg",
            "GERENTE",
            "TURNO",
            "LOCAÇÃO",
            "PESO",
            "Qtd. M.O",
            "Fator ponderado",
            "Target Mensal",
            "Target Week 1",
            "Realizado",
            "Target Week 2",
            "Realizado2",
            "Target Week 3",
            "Realizado3",
            "Target Week 4",
            "Realizado4",
        ]
        for index, value in enumerate(header, start=1):
            worksheet.cell(row=4, column=index, value=value)
        worksheet.append([])
        worksheet.append([
            1, "GERENTE A", "1 TURNO", None, 1, 10, "=E6*F6", "=Z2",
            5, 3, 4, 2, 3, 1, 2, 0,
        ])
        worksheet.append([
            2, "GERENTE B", "2 TURNO", None, 3, 10, "=E7*F7", "=Z2",
            15, 5, 10, 3, 8, 2, 7, 1,
        ])
        worksheet.append(["Total programado mês"])
        for index, value in enumerate(header, start=1):
            worksheet.cell(row=10, column=index, value=value)
        worksheet.append([
            3, "LIDER A", "1 TURNO", None, 1, 10, "=E11*F11", "=Z2",
            5, 2, 4, 2, 3, 1, 2, 1,
        ])
        worksheet.append(["Total programado mês"])
        buffer = BytesIO()
        workbook.save(buffer)
        arquivo = SimpleNamespace(getvalue=buffer.getvalue)

        resultado = ler_target_gerencia_excel(arquivo, date(2026, 10, 20))

        self.assertEqual(resultado["meta_mensal"], 100)
        self.assertEqual(len(resultado["grupos"]), 2)
        gerentes = resultado["grupos"][0]["registros"]
        lideres = resultado["grupos"][1]["registros"]
        self.assertEqual(gerentes["Target Mensal Calculado"].tolist(), [25, 75])
        self.assertEqual(gerentes["Total Realizado Calculado"].tolist(), [6, 11])
        self.assertEqual(lideres["Target Mensal Calculado"].tolist(), [25])
        self.assertEqual(lideres["Total Realizado Calculado"].tolist(), [6])


if __name__ == "__main__":
    unittest.main()
