from datetime import date, datetime, timedelta
from io import BytesIO
import json
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

import openpyxl
import pandas as pd
from sqlalchemy import func, inspect, select

from core.cn_management import (
    _contagens_status,
    _codigo_linha_para_familia,
    _dados_gantt_cns,
    _figura_gantt_cns,
    _formatar_identificador_cn,
    _garantir_responsavel_usuario,
    _indicadores_dashboard,
    _referencias_para_familia,
    analisar_planilha_cns,
)
from database.calibration_store import (
    DATABASE_URL_SECRET,
    REQUIRE_SHARED_DATABASE_SECRET,
    _engine_para_url,
    autenticar_usuario,
    aprovar_solicitacao,
    criar_usuario,
    provisionar_administrador,
)
from database.cn_store import (
    ETAPAS_CN,
    FAMILIA_REFERENCIA_GLOBAL,
    administrar_responsavel,
    atualizar_etapa,
    carregar_auditoria,
    carregar_auditoria_administracao,
    baixar_backup_cns,
    cn_auditoria,
    cn_notificacoes,
    cn_registros,
    carregar_cns,
    carregar_responsaveis,
    carregar_notificacoes,
    carregar_referencias_cns,
    carregar_todas_notificacoes_administracao,
    confirmar_recebimento,
    confirmar_recebimento_etapa,
    criar_cn,
    criar_backup_cns,
    data_para_iso,
    excluir_cn,
    importar_cns,
    inicializar_banco,
    limpar_cns_importadas,
    limpar_dados_cns,
    listar_backups_cns,
    listar_numeros_cns_importadas,
    listar_responsaveis_administracao,
    processar_escalonamentos,
    registrar_alerta_importacao,
    reimportar_base_cns,
    salvar_referencias_cns,
    salvar_responsavel,
    resumir_cns,
    sincronizar_status_cns,
)
from database.calibration_store import _engine


def criar_planilha_cns(valor_etapa: object = datetime(2026, 10, 18)) -> bytes:
    workbook = openpyxl.Workbook()
    controle = workbook.active
    controle.title = "CONTROLE DE CN'S 2026"
    controle.append([])
    controle.append(
        [
            "Contador",
            "Família",
            "Descrição",
            "Código da linha",
            "CN'S",
            "Resumo",
            *ETAPAS_CN,
            "Observações",
            "Status",
        ]
    )
    controle.append(
        [
            7,
            "Família A",
            "Descrição A",
            "Linha antiga",
            "CN-700",
            "Resumo da alteração",
            valor_etapa,
            "NA",
            " - ",
            "_",
            "?",
            "NULL",
            None,
            None,
            None,
            None,
            "Observação importada",
            "IMPLEMENTADO",
        ]
    )
    apoio = workbook.create_sheet("Planilha1")
    apoio.append(["Família A", None, "Descrição A", "Linha automática"])
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


class CNManagementTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.database_url = (
            "sqlite:///" + (Path(self.temp_dir.name) / "cn.db").as_posix()
        )
        self.env_patch = patch.dict(
            os.environ,
            {
                DATABASE_URL_SECRET: self.database_url,
                REQUIRE_SHARED_DATABASE_SECRET: "false",
            },
        )
        self.env_patch.start()
        inicializar_banco(self.database_url)
        provisionar_administrador(
            "Admin CN",
            "admin@example.com",
            "SenhaSegura!123",
            self.database_url,
        )
        salvar_responsavel(
            "Admin CN",
            "admin@example.com",
            "Qualidade",
            "admin@example.com",
            self.database_url,
        )

    def tearDown(self) -> None:
        self.env_patch.stop()
        _engine_para_url(self.database_url).dispose()
        _engine_para_url.cache_clear()
        self.temp_dir.cleanup()

    def test_import_reads_reference_line_and_never_marks_plan_as_completion(self):
        registros, alertas = analisar_planilha_cns(criar_planilha_cns())
        self.assertEqual(alertas, [])
        self.assertEqual(len(registros), 1)
        self.assertEqual(registros[0]["numero_cn"], "7")
        self.assertEqual(registros[0]["codigo_cn"], "CN-700")
        self.assertEqual(registros[0]["codigo_linha"], "Linha automática")
        self.assertEqual(registros[0]["etapas"][ETAPAS_CN[0]], "2026-10-18")
        self.assertEqual(registros[0]["etapas"][ETAPAS_CN[1]], "NA")

        registros[0]["responsavel_email"] = "admin@example.com"
        resultado = importar_cns(registros, "admin@example.com", self.database_url)
        self.assertEqual(resultado, {"novas": 1, "duplicadas": 0})
        cns = carregar_cns(self.database_url)
        self.assertEqual(cns.iloc[0]["status"], "ATRASADO")
        self.assertIsNone(cns.iloc[0]["etapas"][0]["data_conclusao"])
        self.assertEqual(cns.iloc[0]["etapas"][0]["status_calculado"], "ATRASADO")

    def test_cn_summary_initializes_empty_database(self):
        with tempfile.TemporaryDirectory() as diretorio:
            database_url = "sqlite:///" + (Path(diretorio) / "cn-summary.db").as_posix()
            try:
                self.assertEqual(resumir_cns(database_url), (0, 0))
            finally:
                _engine_para_url(database_url).dispose()

    def test_approved_user_is_added_as_responsible_when_none_are_active(self):
        responsavel = carregar_responsaveis(self.database_url).iloc[0]
        administrar_responsavel(
            int(responsavel["id"]),
            responsavel["usuario"],
            responsavel["email"],
            responsavel["setor"],
            False,
            "admin@example.com",
            self.database_url,
        )

        responsaveis = _garantir_responsavel_usuario(
            {
                "nome": "Admin CN",
                "email": "admin@example.com",
                "perfil": "ADMIN",
            },
            self.database_url,
        )

        self.assertEqual(responsaveis["email"].tolist(), ["admin@example.com"])
        self.assertEqual(responsaveis.iloc[0]["setor"], "Não informado")

    def test_active_cn_users_are_added_as_responsibles_and_pending_users_are_omitted(self):
        request_id = criar_usuario(
            "Giovana Oliveira",
            "giovana@example.com",
            "SenhaSegura!456",
        )
        usuario_sessao = {
            "nome": "Admin CN",
            "email": "admin@example.com",
            "perfil": "ADMIN",
        }

        responsaveis = _garantir_responsavel_usuario(
            usuario_sessao,
            self.database_url,
        )
        self.assertNotIn("giovana@example.com", responsaveis["email"].tolist())

        aprovar_solicitacao(request_id, "admin@example.com", "QUALIDADE")
        responsaveis = _garantir_responsavel_usuario(
            usuario_sessao,
            self.database_url,
        )

        giovana = responsaveis.loc[
            responsaveis["email"] == "giovana@example.com"
        ].iloc[0]
        self.assertEqual(giovana["usuario"], "Giovana Oliveira")
        self.assertEqual(giovana["setor"], "Não informado")

        administrar_responsavel(
            int(giovana["id"]),
            giovana["usuario"],
            giovana["email"],
            "Engenharia",
            True,
            "admin@example.com",
            self.database_url,
        )
        responsaveis = _garantir_responsavel_usuario(
            usuario_sessao,
            self.database_url,
        )
        giovana = responsaveis.loc[
            responsaveis["email"] == "giovana@example.com"
        ]
        self.assertEqual(len(giovana), 1)
        self.assertEqual(giovana.iloc[0]["setor"], "Engenharia")

    def test_leitor_is_not_added_as_a_responsible_for_cn_registration(self):
        request_id = criar_usuario(
            "Leitor CN",
            "leitor@example.com",
            "SenhaSegura!456",
        )
        aprovar_solicitacao(request_id, "admin@example.com", "LEITOR")

        responsaveis = _garantir_responsavel_usuario(
            {
                "nome": "Admin CN",
                "email": "admin@example.com",
                "perfil": "ADMIN",
            },
            self.database_url,
        )

        self.assertNotIn("leitor@example.com", responsaveis["email"].tolist())

    def test_import_matches_line_code_by_family_and_description(self):
        workbook = openpyxl.load_workbook(BytesIO(criar_planilha_cns()))
        workbook["CONTROLE DE CN'S 2026"]["B3"] = "Família B"
        workbook["Planilha1"].append(
            ["Família B", None, "Descrição A", "Linha B"]
        )
        buffer = BytesIO()
        workbook.save(buffer)
        workbook.close()

        registros, _ = analisar_planilha_cns(buffer.getvalue())

        self.assertEqual(registros[0]["codigo_linha"], "Linha B")

    def test_reference_lists_persist_for_family_description_and_line_code(self):
        referencias = {
            "linhas": [
                {
                    "familia": "Família A",
                    "descricao": "Descrição A",
                    "codigo_linha": "Linha automática",
                },
                {
                    "familia": "Família B",
                    "descricao": "Descrição B",
                    "codigo_linha": "Linha B",
                },
                {
                    "familia": float("nan"),
                    "descricao": float("nan"),
                    "codigo_linha": float("nan"),
                },
            ]
        }
        salvar_referencias_cns(
            referencias,
            "admin@example.com",
            self.database_url,
        )

        carregadas = carregar_referencias_cns(self.database_url)
        self.assertEqual(carregadas["familias"], ["Família A", "Família B"])
        self.assertEqual(
            carregadas["descricoes"],
            ["Descrição A", "Descrição B"],
        )
        self.assertEqual(
            carregadas["codigo_linha"]["descrição a"],
            "Linha automática",
        )
        self.assertEqual(carregadas["linhas"], referencias["linhas"][:2])

    def test_initial_master_catalog_is_seeded_once_with_family_description_and_line(self):
        referencias = carregar_referencias_cns(self.database_url)
        esperado = {
            ("GM BSUV", "BODY BSUV"): "GM GEM-20M",
            ("GM BSUV", "CONSOLE BSUV"): "GM GEM-80M",
            ("GM BSUV", "FOG LAMP BSUV"): "GM GEM-60M",
            ("GM BSUV", "FRONT BUMPER BSUV"): "GM GEM-90M",
            ("GM BSUV", "FRONT DOOR CO DRIVER BSUV"): "GM BSUV-PKP-70M",
            ("GM BSUV", "FRONT DOOR DRIVER BSUV"): "GM BSUV-PKP-74M",
            ("GM BSUV", "IP EXTENSION BSUV"): "GM GEM-42M",
            ("GM BSUV", "LIFT GATE JPR BSUV"): "GM GEM-43M",
            ("GM BSUV", "LIFT GATE WRG BSUV"): "GM GEM-41M",
            ("GM BSUV", "REAR BUMPER BSUV"): "GM GEM-91M",
            ("GM BSUV", "REAR DOOR LEFT BSUV"): "GM GEM-75M",
            ("GM BSUV", "REAR DOOR RIGHT BSUV"): "GM GEM-71M",
            ("GM BSUV", "ROOF BSUV"): "GM GEM-50M",
            ("GM BSUV", "SUN ROOF BSUV"): "GM GEM-50M",
            ("GM PKP", "BODY PKP"): "GM PKP-21M",
            ("GM PKP", "BODY REAR JPR PKP"): "GM PKP-110M",
            ("GM PKP", "BODY REAR WRG PKP"): "GM PKP-121M",
            ("GM PKP", "CARGO LAMP PKP"): "GM PKP-40M",
            ("GM PKP", "CONSOLE PKP"): "GM PKP-81M",
            ("GM PKP", "END GATE PKP"): "GM PKP-100M",
            ("GM PKP", "FOG LAMP PKP"): "GM PKP-60M",
            ("GM PKP", "FRONT BUMPER PKP"): "GM PKP-60/65M",
            ("GM PKP", "FRONT DOOR CO DRIVER PKP"): "GM BSUV-PKP-70M",
            ("GM PKP", "FRONT DOOR DRIVER PKP"): "GM BSUV-PKP-74M",
            ("GM PKP", "IP EXTENSION PKP"): "GM PKP-42M",
            ("GM PKP", "REAR BUMPER PKP"): "GM PKP-120M",
            ("GM PKP", "REAR DOOR LEFT PKP"): "GM PKP-73M",
            ("GM PKP", "REAR DOOR RIGHT PKP"): "GM PKP-72M",
            ("GM PKP", "REAR VIEW CAMERA PKP"): "GM PKP-40M",
            ("GM PKP", "ROOF PKP"): "GM PKP-51M",
            ("GM BSUV/PKP", "CONSOLE BSUV/PKP"): "GM BSUV/PKP-82M BC",
            ("GM BSUV/PKP", "ENGINE LINHA 1 - MM1"): "GM BSUV/PKP-34M",
            ("GM BSUV/PKP", "ENGINE LINHA 2 - MM1"): "GM BSUV/PKP-30/32M",
            ("GM BSUV/PKP", "ENGINE LINHA 3 - MM3"): "GM BSUV/PKP-31/33M",
            ("GM BSUV/PKP", "IP BSUV/PKP"): "GM GEM-12M BC",
            ("VW", "CABO MASSA"): "NF-91M/E",
            ("VW", "CABO BATERIA"): "FX-92M/E",
            ("VW", "AMAROK MIRGOR"): "AMK-80M BC/AMK-81M BC",
            ("VW", "YAPP"): "VW YAPP-15M",
            ("VW", "PRINCIPAL GP"): "VW GP -13M",
            ("VW", "PRINCIPAL KSK 246"): "VW 246 - 20M BC",
            ("GM BSUV/PKP", "ENGINE"): "GM BSUV/PKP-36M BC",
            ("GM PKP", "ENGINE MM3"): "GM PKP-35M",
            ("GM BSUV/PKP", "FRONT BUMPER PKP-BSUV"): "GM BSUV-PKP-60E",
            (FAMILIA_REFERENCIA_GLOBAL, "GERAL"): "afeta todas as familias",
        }
        self.assertEqual(
            {
                (linha["familia"], linha["descricao"]): linha["codigo_linha"]
                for linha in referencias["linhas"]
            },
            esperado,
        )
        referencias_gm = _referencias_para_familia(
            referencias["linhas"],
            "GM PKP",
        )
        self.assertIn(
            {
                "familia": FAMILIA_REFERENCIA_GLOBAL,
                "descricao": "GERAL",
                "codigo_linha": "afeta todas as familias",
            },
            referencias_gm,
        )
        self.assertIn(
            {
                "familia": "GM PKP",
                "descricao": "ENGINE MM3",
                "codigo_linha": "GM PKP-35M",
            },
            referencias_gm,
        )
        self.assertEqual(
            _codigo_linha_para_familia(referencias_gm, "GM PKP", "GERAL"),
            "afeta todas as familias",
        )

        linhas_antes = referencias["linhas"]
        inicializar_banco(self.database_url)
        self.assertEqual(
            carregar_referencias_cns(self.database_url)["linhas"],
            linhas_antes,
        )

    def test_quality_user_updates_master_data_shared_with_other_users(self):
        solicitacao_id = criar_usuario(
            "Pessoa Qualidade",
            "qualidade@example.com",
            "SenhaSegura!456",
        )
        aprovar_solicitacao(
            solicitacao_id,
            "admin@example.com",
            "QUALIDADE",
        )
        referencias = {
            "linhas": [
                {
                    "familia": "Família Compartilhada",
                    "descricao": "Descrição Compartilhada",
                    "codigo_linha": "Linha Compartilhada",
                }
            ]
        }

        salvar_referencias_cns(
            referencias,
            "qualidade@example.com",
            self.database_url,
        )

        self.assertEqual(
            carregar_referencias_cns(self.database_url)["linhas"],
            referencias["linhas"],
        )

    def test_approved_users_share_cn_data_after_database_reconnect(self):
        solicitacao_id = criar_usuario(
            "Nickolas",
            "nickolas@example.com",
            "SenhaSegura!456",
        )
        aprovar_solicitacao(
            solicitacao_id,
            "admin@example.com",
            "QUALIDADE",
        )
        solicitacao_id = criar_usuario(
            "Giovana",
            "giovana@example.com",
            "SenhaSegura!789",
        )
        aprovar_solicitacao(
            solicitacao_id,
            "admin@example.com",
            "LEITOR",
        )

        etapas = {
            nome: {
                "data_planejada": "NA",
                "responsavel_email": "admin@example.com",
            }
            for nome in ETAPAS_CN
        }
        criar_cn(
            {
                "numero_cn": "CN-SHARED",
                "codigo_cn": "CODE-SHARED",
                "familia": "Família Compartilhada",
                "descricao": "Descrição Compartilhada",
                "codigo_linha": "Linha Compartilhada",
                "resumo_modificacao": "Criada por Nickolas",
                "observacao": "Dados globais da CN",
                "responsavel_email": "admin@example.com",
                "criticidade": "Alta",
            },
            etapas,
            "nickolas@example.com",
            self.database_url,
        )

        _engine_para_url(self.database_url).dispose()
        giovana = autenticar_usuario(
            "giovana@example.com",
            "SenhaSegura!789",
        )
        self.assertIsNotNone(giovana)
        cns = carregar_cns(self.database_url)
        cn = cns.loc[cns["numero_cn"] == "CN-SHARED"].iloc[0]

        self.assertEqual(cn["resumo_modificacao"], "Criada por Nickolas")
        self.assertEqual(cn["observacao"], "Dados globais da CN")
        self.assertEqual(len(cn["etapas"]), len(ETAPAS_CN))
        self.assertEqual(cn["etapas"][0]["etapa"], ETAPAS_CN[0])

    def test_multiple_dates_raise_auditable_import_warning(self):
        registros, alertas = analisar_planilha_cns(
            criar_planilha_cns("10/10/2026\n15/10/2026")
        )
        self.assertEqual(len(registros), 1)
        self.assertEqual(registros[0]["etapas"][ETAPAS_CN[0]], "NA")
        self.assertEqual(len(alertas), 1)
        self.assertEqual(alertas[0]["coluna"], "G")
        self.assertEqual(alertas[0]["linha"], 3)
        registrar_alerta_importacao(
            "admin@example.com",
            "7",
            "G",
            3,
            alertas[0]["detalhe"],
            self.database_url,
        )
        auditoria = carregar_auditoria(self.database_url)
        self.assertEqual(auditoria.iloc[0]["CN"], "7")
        self.assertEqual(auditoria.iloc[0]["Campo"], "G")
        self.assertEqual(auditoria.iloc[0]["Tipo Alteração"], "ALERTA DE IMPORTAÇÃO")

    def test_create_cn_starts_with_pending_confirmation_and_notification(self):
        etapas = {
            nome: {
                "data_planejada": (
                    date.today() + timedelta(days=5) if nome == ETAPAS_CN[0] else "NA"
                ),
                "responsavel_email": "admin@example.com",
            }
            for nome in ETAPAS_CN
        }
        cn_id = criar_cn(
            {
                "numero_cn": "CN-001",
                "codigo_cn": "CODE-001",
                "familia": "Família A",
                "descricao": "Descrição A",
                "codigo_linha": "Linha A",
                "resumo_modificacao": "Mudança de teste",
                "observacao": "",
                "responsavel_email": "admin@example.com",
                "criticidade": "Alta",
            },
            etapas,
            "admin@example.com",
            self.database_url,
        )
        cn = carregar_cns(self.database_url).iloc[0]
        self.assertEqual(int(cn["id"]), cn_id)
        self.assertEqual(cn["status"], "PROGRAMADO")
        self.assertEqual(cn["status_confirmacao"], "PENDENTE")
        self.assertEqual(
            carregar_notificacoes("admin@example.com", self.database_url).iloc[0]["Tipo"],
            "CN Criada",
        )

    def test_existing_cn_database_migrates_stage_receipt_columns(self):
        engine = _engine_para_url(self.database_url)
        with engine.begin() as conexao:
            conexao.exec_driver_sql(
                "ALTER TABLE cn_etapas DROP COLUMN data_recebimento"
            )
            conexao.exec_driver_sql(
                "ALTER TABLE cn_etapas DROP COLUMN recebido_por"
            )

        inicializar_banco(self.database_url)

        colunas = {
            coluna["name"]
            for coluna in inspect(engine).get_columns("cn_etapas")
        }
        self.assertIn("data_recebimento", colunas)
        self.assertIn("recebido_por", colunas)

    def test_cn_summary_is_read_only_and_uses_calculated_stage_status(self):
        etapas = {
            nome: {
                "data_planejada": (
                    date.today() - timedelta(days=1)
                    if nome == ETAPAS_CN[0]
                    else "NA"
                ),
                "responsavel_email": "admin@example.com",
            }
            for nome in ETAPAS_CN
        }
        criar_cn(
            {
                "numero_cn": "CN-002",
                "codigo_cn": "CODE-002",
                "familia": "Família A",
                "descricao": "Descrição A",
                "codigo_linha": "Linha A",
                "resumo_modificacao": "Mudança de teste",
                "observacao": "",
                "responsavel_email": "admin@example.com",
                "criticidade": "Alta",
            },
            etapas,
            "admin@example.com",
            self.database_url,
        )
        auditoria_antes = len(carregar_auditoria(self.database_url))
        notificacoes_antes = len(
            carregar_notificacoes("admin@example.com", self.database_url)
        )

        self.assertEqual(resumir_cns(self.database_url), (1, 1))

        self.assertEqual(len(carregar_auditoria(self.database_url)), auditoria_antes)
        self.assertEqual(
            len(carregar_notificacoes("admin@example.com", self.database_url)),
            notificacoes_antes,
        )

    def test_loading_cns_calculates_status_without_writing_records_audit_or_notifications(self):
        etapas = {
            nome: {
                "data_planejada": (
                    date.today() - timedelta(days=1)
                    if nome == ETAPAS_CN[0]
                    else "NA"
                ),
                "responsavel_email": "admin@example.com",
            }
            for nome in ETAPAS_CN
        }
        cn_id = criar_cn(
            {
                "numero_cn": "CN-READ-ONLY",
                "codigo_cn": "CODE-READ-ONLY",
                "familia": "Família A",
                "descricao": "CN para teste de leitura",
                "codigo_linha": "Linha A",
                "resumo_modificacao": "Verificar consulta sem gravação",
                "observacao": "",
                "responsavel_email": "admin@example.com",
                "criticidade": "Alta",
            },
            etapas,
            "admin@example.com",
            self.database_url,
        )
        engine = _engine(self.database_url)
        with engine.connect() as conexao:
            status_antes = conexao.execute(
                select(cn_registros.c.status).where(cn_registros.c.id == cn_id)
            ).scalar_one()
            auditorias_antes = conexao.execute(
                select(func.count()).select_from(cn_auditoria)
            ).scalar_one()
            notificacoes_antes = conexao.execute(
                select(func.count()).select_from(cn_notificacoes)
            ).scalar_one()

        cns = carregar_cns(self.database_url)

        self.assertEqual(cns.iloc[0]["status"], "ATRASADO")
        with engine.connect() as conexao:
            self.assertEqual(
                conexao.execute(
                    select(cn_registros.c.status).where(cn_registros.c.id == cn_id)
                ).scalar_one(),
                status_antes,
            )
            self.assertEqual(
                conexao.execute(
                    select(func.count()).select_from(cn_auditoria)
                ).scalar_one(),
                auditorias_antes,
            )
            self.assertEqual(
                conexao.execute(
                    select(func.count()).select_from(cn_notificacoes)
                ).scalar_one(),
                notificacoes_antes,
            )

    def test_status_sync_persists_calculated_changes_and_audits_them(self):
        etapas = {
            nome: {
                "data_planejada": (
                    date.today() - timedelta(days=1)
                    if nome == ETAPAS_CN[0]
                    else "NA"
                ),
                "responsavel_email": "admin@example.com",
            }
            for nome in ETAPAS_CN
        }
        cn_id = criar_cn(
            {
                "numero_cn": "CN-STATUS-SYNC",
                "codigo_cn": "CODE-STATUS-SYNC",
                "familia": "Família A",
                "descricao": "CN para sincronização explícita",
                "codigo_linha": "Linha A",
                "resumo_modificacao": "Verificar sincronização explícita",
                "observacao": "",
                "responsavel_email": "admin@example.com",
                "criticidade": "Alta",
            },
            etapas,
            "admin@example.com",
            self.database_url,
        )

        self.assertEqual(sincronizar_status_cns(self.database_url), 1)
        with _engine(self.database_url).connect() as conexao:
            self.assertEqual(
                conexao.execute(
                    select(cn_registros.c.status).where(cn_registros.c.id == cn_id)
                ).scalar_one(),
                "ATRASADO",
            )
        auditoria = carregar_auditoria(self.database_url)
        self.assertIn("ALTERAÇÃO DE STATUS", auditoria["Tipo Alteração"].tolist())

    def test_admin_management_updates_responsibles_and_preserves_delete_audit(self):
        etapas = {
            nome: {
                "data_planejada": "NA",
                "responsavel_email": "admin@example.com",
            }
            for nome in ETAPAS_CN
        }
        cn_id = criar_cn(
            {
                "numero_cn": "CN-DELETE",
                "codigo_cn": "CODE-DELETE",
                "familia": "Família A",
                "descricao": "Descrição a excluir",
                "codigo_linha": "Linha A",
                "resumo_modificacao": "CN de teste administrativo",
                "observacao": "",
                "responsavel_email": "admin@example.com",
                "criticidade": "Média",
            },
            etapas,
            "admin@example.com",
            self.database_url,
        )
        responsaveis = listar_responsaveis_administracao(
            "admin@example.com",
            self.database_url,
        )
        responsavel = responsaveis.iloc[0]
        administrar_responsavel(
            int(responsavel["id"]),
            responsavel["usuario"],
            responsavel["email"],
            "Engenharia",
            True,
            "admin@example.com",
            self.database_url,
        )
        self.assertEqual(
            listar_responsaveis_administracao(
                "admin@example.com",
                self.database_url,
            ).iloc[0]["setor"],
            "Engenharia",
        )

        notificacoes = carregar_todas_notificacoes_administracao(
            "admin@example.com",
            self.database_url,
        )
        self.assertIn("admin@example.com", notificacoes["Destinatário"].tolist())
        excluir_cn(cn_id, "admin@example.com", self.database_url)

        self.assertTrue(carregar_cns(self.database_url).empty)
        auditoria = carregar_auditoria_administracao(
            "admin@example.com",
            self.database_url,
        )
        exclusao = auditoria[auditoria["Tipo Alteração"] == "EXCLUSÃO"]
        self.assertEqual(len(exclusao), 1)
        self.assertEqual(exclusao.iloc[0]["CN"], "CN-DELETE")
        self.assertIn("Descrição a excluir", exclusao.iloc[0]["Valor Antigo"])
        self.assertTrue(
            carregar_todas_notificacoes_administracao(
                "admin@example.com",
                self.database_url,
            ).empty
        )

    def test_full_cleanup_backs_up_cn_data_and_preserves_users_master_and_audit(self):
        registros, _ = analisar_planilha_cns(criar_planilha_cns())
        registros[0]["responsavel_email"] = "admin@example.com"
        importar_cns(registros, "admin@example.com", self.database_url)
        salvar_referencias_cns(
            {
                "linhas": [
                    {
                        "familia": "Família A",
                        "descricao": "Descrição A",
                        "codigo_linha": "Linha automática",
                    }
                ]
            },
            "admin@example.com",
            self.database_url,
        )

        resultado = limpar_dados_cns("admin@example.com", self.database_url)

        self.assertEqual(resultado["cns_removidas"], 1)
        self.assertEqual(resultado["cns_no_backup"], 1)
        self.assertTrue(carregar_cns(self.database_url).empty)
        self.assertTrue(
            carregar_todas_notificacoes_administracao(
                "admin@example.com",
                self.database_url,
            ).empty
        )
        self.assertEqual(
            listar_responsaveis_administracao(
                "admin@example.com",
                self.database_url,
            )["email"].tolist(),
            ["admin@example.com"],
        )
        self.assertEqual(
            carregar_referencias_cns(self.database_url)["familias"],
            ["Família A"],
        )
        self.assertEqual(
            carregar_auditoria_administracao(
                "admin@example.com",
                self.database_url,
            )["Tipo Alteração"].tolist().count("LIMPEZA TOTAL"),
            1,
        )
        backup = json.loads(
            baixar_backup_cns(
                resultado["backup_id"],
                "admin@example.com",
                self.database_url,
            )
        )
        self.assertEqual(backup["cn_registros"][0]["numero_cn"], "7")
        self.assertEqual(len(backup["cn_etapas"]), len(ETAPAS_CN))
        self.assertEqual(
            len(listar_backups_cns("admin@example.com", self.database_url)),
            1,
        )

    def test_reimport_replaces_cn_base_atomically_and_generates_backup(self):
        antigos, _ = analisar_planilha_cns(criar_planilha_cns())
        antigos[0]["responsavel_email"] = "admin@example.com"
        importar_cns(antigos, "admin@example.com", self.database_url)
        novos, _ = analisar_planilha_cns(criar_planilha_cns())
        novos[0]["numero_cn"] = "CN-NOVA"
        novos[0]["codigo_cn"] = "CODE-NOVA"
        novos[0]["responsavel_email"] = "admin@example.com"

        resultado = reimportar_base_cns(
            novos,
            "admin@example.com",
            self.database_url,
        )

        self.assertEqual(resultado["backup_id"], 1)
        self.assertEqual(resultado["cns_removidas"], 1)
        self.assertEqual(resultado["novas"], 1)
        self.assertEqual(carregar_cns(self.database_url)["numero_cn"].tolist(), ["CN-NOVA"])
        backup = json.loads(
            baixar_backup_cns(
                resultado["backup_id"],
                "admin@example.com",
                self.database_url,
            )
        )
        self.assertEqual(backup["cn_registros"][0]["numero_cn"], "7")

    def test_failed_reimport_rolls_back_cleanup_and_backup(self):
        antigos, _ = analisar_planilha_cns(criar_planilha_cns())
        antigos[0]["responsavel_email"] = "admin@example.com"
        importar_cns(antigos, "admin@example.com", self.database_url)
        referencias_antigas = {
            "linhas": [
                {
                    "familia": "Família Antiga",
                    "descricao": "Descrição Antiga",
                    "codigo_linha": "Linha Antiga",
                }
            ]
        }
        salvar_referencias_cns(
            referencias_antigas,
            "admin@example.com",
            self.database_url,
        )
        invalidos, _ = analisar_planilha_cns(criar_planilha_cns())
        invalidos[0]["numero_cn"] = "CN-INVALIDA"
        invalidos[0]["responsavel_email"] = "inexistente@example.com"
        referencias_novas = {
            "linhas": [
                {
                    "familia": "Família Nova",
                    "descricao": "Descrição Nova",
                    "codigo_linha": "Linha Nova",
                }
            ]
        }

        with self.assertRaisesRegex(ValueError, "responsável ativo"):
            reimportar_base_cns(
                invalidos,
                "admin@example.com",
                self.database_url,
                referencias_novas,
            )

        self.assertEqual(carregar_cns(self.database_url)["numero_cn"].tolist(), ["7"])
        self.assertEqual(
            carregar_referencias_cns(self.database_url)["linhas"],
            referencias_antigas["linhas"],
        )
        self.assertTrue(
            listar_backups_cns("admin@example.com", self.database_url).empty
        )

    def test_admin_cleans_only_imported_cns_and_preserves_audit(self):
        registros, alertas = analisar_planilha_cns(criar_planilha_cns())
        self.assertEqual(alertas, [])
        registros[0]["responsavel_email"] = "admin@example.com"
        importar_cns(registros, "admin@example.com", self.database_url)
        etapas = {
            nome: {
                "data_planejada": "NA",
                "responsavel_email": "admin@example.com",
            }
            for nome in ETAPAS_CN
        }
        criar_cn(
            {
                "numero_cn": "CN-MANUAL",
                "codigo_cn": "CODE-MANUAL",
                "familia": "Família A",
                "descricao": "CN manual",
                "codigo_linha": "Linha A",
                "resumo_modificacao": "Cadastro manual",
                "observacao": "",
                "responsavel_email": "admin@example.com",
                "criticidade": "Média",
            },
            etapas,
            "admin@example.com",
            self.database_url,
        )

        self.assertEqual(
            listar_numeros_cns_importadas(
                "admin@example.com",
                self.database_url,
            ),
            ["7"],
        )
        self.assertEqual(
            limpar_cns_importadas("admin@example.com", self.database_url),
            1,
        )

        cns_restantes = carregar_cns(self.database_url)
        self.assertEqual(cns_restantes["numero_cn"].tolist(), ["CN-MANUAL"])
        notificacoes = carregar_todas_notificacoes_administracao(
            "admin@example.com",
            self.database_url,
        )
        self.assertEqual(notificacoes["CN"].tolist(), ["CN-MANUAL"])
        auditoria = carregar_auditoria_administracao(
            "admin@example.com",
            self.database_url,
        )
        self.assertEqual(
            len(
                auditoria[
                    (auditoria["CN"] == "7")
                    & (auditoria["Campo"] == "Importação da planilha")
                    & (auditoria["Tipo Alteração"] == "CRIAÇÃO")
                ]
            ),
            1,
        )
        exclusoes = auditoria[
            auditoria["Tipo Alteração"] == "LIMPEZA DE IMPORTAÇÃO"
        ]
        self.assertEqual(len(exclusoes), 1)
        self.assertEqual(exclusoes.iloc[0]["CN"], "7")
        self.assertIn("CN-700", exclusoes.iloc[0]["Valor Antigo"])

    def test_cleanup_does_not_remove_manually_recreated_imported_number(self):
        registros, alertas = analisar_planilha_cns(criar_planilha_cns())
        self.assertEqual(alertas, [])
        registros[0]["responsavel_email"] = "admin@example.com"
        cn_id_importada = importar_cns(
            registros,
            "admin@example.com",
            self.database_url,
        )
        self.assertEqual(cn_id_importada, {"novas": 1, "duplicadas": 0})
        cn_importada = carregar_cns(self.database_url).iloc[0]
        excluir_cn(
            int(cn_importada["id"]),
            "admin@example.com",
            self.database_url,
        )
        etapas = {
            nome: {
                "data_planejada": "NA",
                "responsavel_email": "admin@example.com",
            }
            for nome in ETAPAS_CN
        }
        criar_cn(
            {
                "numero_cn": "7",
                "codigo_cn": "CODE-RECREATED",
                "familia": "Família A",
                "descricao": "CN recriada manualmente",
                "codigo_linha": "Linha A",
                "resumo_modificacao": "Cadastro manual",
                "observacao": "",
                "responsavel_email": "admin@example.com",
                "criticidade": "Média",
            },
            etapas,
            "admin@example.com",
            self.database_url,
        )

        self.assertEqual(
            listar_numeros_cns_importadas(
                "admin@example.com",
                self.database_url,
            ),
            [],
        )
        self.assertEqual(
            limpar_cns_importadas("admin@example.com", self.database_url),
            0,
        )
        self.assertEqual(carregar_cns(self.database_url).iloc[0]["descricao"], "CN recriada manualmente")

    def test_admin_only_cn_management_endpoints_reject_non_admin_users(self):
        solicitacao_id = criar_usuario(
            "Pessoa Qualidade",
            "qualidade@example.com",
            "SenhaSegura!456",
        )
        aprovar_solicitacao(
            solicitacao_id,
            "admin@example.com",
            "QUALIDADE",
        )
        with self.assertRaises(PermissionError):
            listar_responsaveis_administracao(
                "qualidade@example.com",
                self.database_url,
            )
        with self.assertRaises(PermissionError):
            carregar_auditoria_administracao(
                "qualidade@example.com",
                self.database_url,
            )
        with self.assertRaises(PermissionError):
            carregar_todas_notificacoes_administracao(
                "qualidade@example.com",
                self.database_url,
            )
        with self.assertRaises(PermissionError):
            listar_numeros_cns_importadas(
                "qualidade@example.com",
                self.database_url,
            )
        with self.assertRaises(PermissionError):
            limpar_cns_importadas(
                "qualidade@example.com",
                self.database_url,
            )

    def test_admin_can_create_and_deactivate_cn_responsible(self):
        solicitacao_id = criar_usuario(
            "Pessoa Qualidade",
            "qualidade@example.com",
            "SenhaSegura!456",
        )
        aprovar_solicitacao(
            solicitacao_id,
            "admin@example.com",
            "QUALIDADE",
        )
        administrar_responsavel(
            None,
            "Pessoa Qualidade",
            "qualidade@example.com",
            "Engenharia",
            True,
            "admin@example.com",
            self.database_url,
        )
        responsaveis = listar_responsaveis_administracao(
            "admin@example.com",
            self.database_url,
        )
        novo = responsaveis[
            responsaveis["email"] == "qualidade@example.com"
        ].iloc[0]
        self.assertEqual(int(novo["ativo"]), 1)

        administrar_responsavel(
            int(novo["id"]),
            novo["usuario"],
            novo["email"],
            novo["setor"],
            False,
            "admin@example.com",
            self.database_url,
        )
        responsaveis = listar_responsaveis_administracao(
            "admin@example.com",
            self.database_url,
        )
        novo = responsaveis[
            responsaveis["email"] == "qualidade@example.com"
        ].iloc[0]
        self.assertEqual(int(novo["ativo"]), 0)

    def test_repeated_code_is_allowed_when_reference_number_differs(self):
        registros, _ = analisar_planilha_cns(criar_planilha_cns())
        second = dict(registros[0])
        second["numero_cn"] = "8"
        second["etapas"] = dict(registros[0]["etapas"])
        registros.append(second)
        for registro in registros:
            registro["responsavel_email"] = "admin@example.com"
        resultado = importar_cns(registros, "admin@example.com", self.database_url)
        self.assertEqual(resultado, {"novas": 2, "duplicadas": 0})
        self.assertEqual(len(carregar_cns(self.database_url)), 2)

    def test_date_normalization_and_stage_status(self):
        for value in (None, "", "-", "_", "?", "NULL"):
            self.assertEqual(data_para_iso(value, "G", 3)[0], "NA")
        self.assertEqual(data_para_iso("19/10/2026", "G", 3)[0], "2026-10-19")
        self.assertIsNotNone(data_para_iso("10/10/2026\n15/10/2026", "G", 3)[1])

    def test_display_identifier_uses_sequence_and_pads_three_digits(self):
        self.assertEqual(_formatar_identificador_cn("1"), "CN #001")
        self.assertEqual(_formatar_identificador_cn("125"), "CN #125")
        self.assertEqual(
            _formatar_identificador_cn("3616947-06"),
            "CN #3616947-06",
        )

    def test_dashboard_status_buckets_normalize_text_and_reconcile_total(self):
        dataframe = pd.DataFrame(
            {
                "status": [
                    " ATRASADO ",
                    "PROGRAMADO",
                    "IMPLEMENTADO",
                    "atrasada",
                    None,
                ],
                "status_confirmacao": [
                    "PENDENTE",
                    "CONFIRMADO",
                    "CONFIRMADO",
                    None,
                    "PENDENTE",
                ],
                "criticidade": ["Alta", "Média", "Alta", "Alta", "Baixa"],
            }
        )
        self.assertEqual(
            _contagens_status(dataframe),
            {
                "IMPLEMENTADO": 1,
                "PROGRAMADO": 1,
                "ATRASADO": 2,
                "OUTROS": 1,
            },
        )
        indicadores = _indicadores_dashboard(dataframe)
        self.assertEqual(indicadores["total"], 5)
        self.assertEqual(indicadores["implementadas"], 1)
        self.assertEqual(indicadores["programadas"], 1)
        self.assertEqual(indicadores["atrasadas"], 2)
        self.assertEqual(indicadores["sem_status"], 1)
        self.assertEqual(
            indicadores["implementadas"]
            + indicadores["programadas"]
            + indicadores["atrasadas"]
            + indicadores["sem_status"],
            indicadores["total"],
        )

    def test_confirmation_audit_and_automatic_implementation_status(self):
        registros, _ = analisar_planilha_cns(
            criar_planilha_cns(date.today() + timedelta(days=10))
        )
        registros[0]["responsavel_email"] = "admin@example.com"
        importar_cns(registros, "admin@example.com", self.database_url)
        cn_id = int(carregar_cns(self.database_url).iloc[0]["id"])
        confirmar_recebimento(cn_id, "admin@example.com", self.database_url)
        for etapa in ETAPAS_CN:
            atualizar_etapa(
                cn_id,
                etapa,
                date.today() + timedelta(days=10),
                date.today(),
                "admin@example.com",
                "admin@example.com",
                self.database_url,
            )
        cns = carregar_cns(self.database_url)
        self.assertEqual(cns.iloc[0]["status"], "IMPLEMENTADO")
        self.assertEqual(cns.iloc[0]["status_confirmacao"], "CONFIRMADO")
        self.assertEqual(
            cns.iloc[0]["etapas"][0]["status_calculado"],
            "CONCLUÍDO",
        )
        auditoria = carregar_auditoria(self.database_url)
        self.assertIn("CONFIRMAÇÃO DE RECEBIMENTO", auditoria["Tipo Alteração"].tolist())
        self.assertIn("ALTERAÇÃO DE DATAS", auditoria["Tipo Alteração"].tolist())
        notificacoes = carregar_notificacoes("admin@example.com", self.database_url)
        self.assertIn("CN Criada", notificacoes["Tipo"].tolist())

    def test_stage_receipt_confirms_progress_and_closes_cn_after_all_stages(self):
        etapas = {
            nome: {
                "data_planejada": date.today() + timedelta(days=10),
                "responsavel_email": "admin@example.com",
            }
            for nome in ETAPAS_CN
        }
        cn_id = criar_cn(
            {
                "numero_cn": "CN-ETAPAS",
                "codigo_cn": "CODE-ETAPAS",
                "familia": "Família A",
                "descricao": "CN por etapas",
                "codigo_linha": "Linha A",
                "resumo_modificacao": "Confirmação etapa a etapa",
                "observacao": "",
                "responsavel_email": "admin@example.com",
                "criticidade": "Média",
            },
            etapas,
            "admin@example.com",
            self.database_url,
        )

        confirmar_recebimento_etapa(
            cn_id,
            ETAPAS_CN[0],
            "admin@example.com",
            self.database_url,
        )
        cn = carregar_cns(self.database_url).iloc[0]
        primeira_etapa = next(
            etapa for etapa in cn["etapas"] if etapa["etapa"] == ETAPAS_CN[0]
        )
        self.assertEqual(primeira_etapa["status_calculado"], "CONCLUÍDO")
        self.assertEqual(primeira_etapa["data_recebimento"], date.today().isoformat())
        self.assertEqual(primeira_etapa["data_conclusao"], date.today().isoformat())
        self.assertEqual(cn["status_confirmacao"], "PENDENTE")
        self.assertEqual(cn["status"], "PROGRAMADO")

        for nome_etapa in ETAPAS_CN[1:]:
            confirmar_recebimento_etapa(
                cn_id,
                nome_etapa,
                "admin@example.com",
                self.database_url,
            )
        cn = carregar_cns(self.database_url).iloc[0]
        self.assertEqual(cn["status_confirmacao"], "CONFIRMADO")
        self.assertEqual(cn["status"], "IMPLEMENTADO")
        self.assertTrue(
            all(etapa["data_recebimento"] for etapa in cn["etapas"])
        )
        auditoria = carregar_auditoria(self.database_url)
        self.assertEqual(
            int(auditoria["Tipo Alteração"].eq("CONFIRMAÇÃO DE RECEBIMENTO").sum()),
            len(ETAPAS_CN) + 1,
        )

    def test_gantt_data_uses_schedule_dates_and_filters_by_year(self):
        criada = date(2026, 1, 5)
        dataframe = pd.DataFrame(
            [
                {
                    "numero_cn": "CN-2026",
                    "data_criacao": criada,
                    "familia": "Família A",
                    "etapas": [
                        {
                            "etapa": ETAPAS_CN[0],
                            "data_planejada": "2026-01-15",
                            "data_conclusao": "2026-01-12",
                            "data_recebimento": None,
                            "responsavel_email": "admin@example.com",
                            "status_calculado": "CONCLUÍDO",
                        },
                        {
                            "etapa": ETAPAS_CN[1],
                            "data_planejada": "2026-02-10",
                            "data_conclusao": None,
                            "data_recebimento": None,
                            "responsavel_email": "admin@example.com",
                            "status_calculado": "EM ANDAMENTO",
                        },
                        {
                            "etapa": ETAPAS_CN[2],
                            "data_planejada": "2027-01-10",
                            "data_conclusao": None,
                            "data_recebimento": None,
                            "responsavel_email": "admin@example.com",
                            "status_calculado": "PENDENTE",
                        },
                    ],
                }
            ]
        )

        tarefas = _dados_gantt_cns(dataframe, 2026)

        self.assertEqual(len(tarefas), 2)
        self.assertEqual(tarefas.iloc[0]["Início"], date(2026, 1, 5))
        self.assertEqual(tarefas.iloc[0]["Prazo"], date(2026, 1, 15))
        self.assertEqual(tarefas.iloc[0]["Conclusão real"], date(2026, 1, 12))
        self.assertEqual(tarefas.iloc[1]["Início"], date(2026, 1, 15))
        self.assertEqual(tarefas.iloc[1]["Prazo"], date(2026, 2, 10))
        figura = _figura_gantt_cns(tarefas, 2026)
        self.assertTrue(any(trace.name == "Conclusão real" for trace in figura.data))
        self.assertEqual(figura.layout.xaxis.range[0], date(2026, 1, 1))
        self.assertEqual(figura.layout.xaxis.range[1], date(2027, 1, 1))

    def test_confirmation_escalation_emits_reminder_and_alert_after_72_hours(self):
        registros, _ = analisar_planilha_cns(criar_planilha_cns(None))
        registros[0]["responsavel_email"] = "admin@example.com"
        importar_cns(registros, "admin@example.com", self.database_url)
        with _engine_para_url(self.database_url).begin() as conexao:
            conexao.exec_driver_sql(
                "UPDATE cn_confirmacoes SET solicitada_em = ?",
                (datetime.now() - timedelta(hours=73),),
            )
        self.assertEqual(processar_escalonamentos(self.database_url), 4)
        notificacoes = carregar_notificacoes("admin@example.com", self.database_url)
        self.assertIn("CN sem confirmação — lembrete", notificacoes["Tipo"].tolist())
        self.assertIn("CN sem confirmação — alerta", notificacoes["Tipo"].tolist())


if __name__ == "__main__":
    unittest.main()
