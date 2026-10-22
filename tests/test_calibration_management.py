from datetime import date, datetime
from io import BytesIO
import unittest

import openpyxl
import pandas as pd
from pypdf import PdfReader

from core.calibration_management import (
    ErroEstruturaCalibracao,
    CORES_STATUS_CALIBRACAO,
    _pdf_calibracao,
    _figura_barras_categoricas,
    _figura_vencimentos_mensais,
    _filtrar_status_relatorio,
    carregar_calibracao,
    classificar_vencimento,
    exportar_excel_calibracao,
    PALETA_GRAFICOS_CALIBRACAO,
    normalizar_codigo,
    validar_estrutura_cal,
)


def criar_planilha_calibracao(
    proxima_calibracao="=FORMULA_SE_IGNORADA",
    status_formula="=STATUS_IGNORADO",
    incluir_coluna_equipamento=True,
):
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "CAL"
    worksheet.merge_cells("A1:D1")
    worksheet["A1"] = "Gestão de Calibração"
    headers = [
        "NÚMERO\nEQUIPAMENTO",
        "EQUIPAMENTO",
        "DATA CALIBRAÇÃO",
        "PRÓXIMA CALIBRAÇÃO",
        "LOCAL DE USO",
        "SETOR",
        "STATUS CALIBRAÇÃO",
        "CALIBRAÇÃO",
        "FABRICANTE",
        "RASTREABILIDADE",
        "UNIDADE",
        "CAPACIDADE INSTRUMENTO",
        "VALOR DIVISÃO (RESOLUÇÃO)",
        "FAIXA DE USO",
        "STATUS EQUIPAMENTO",
        "FREQUÊNCIA CALIBRAÇÃO",
    ]
    worksheet.append(headers)
    worksheet.append(
        [
            "E-10060",
            "Paquímetro",
            datetime(2025, 1, 1),
            proxima_calibracao,
            "Laboratório",
            "Qualidade",
            status_formula,
            "Externa",
            "-",
            "-",
            "mm",
            "150",
            "0,01",
            "0 a 150",
            "ATIVO",
            "=FORMULA_FREQUENCIA",
        ]
    )
    if incluir_coluna_equipamento:
        worksheet["B2"] = "EQUIPAMENTO"
    apoio = workbook.create_sheet("DADOS APOIO")
    apoio.append(["Dados de apoio"])
    apoio.append(["EQUIPAMENTO", "VALIDADE CALIBRAÇÃO"])
    apoio.append(["Paquímetro", 365])
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


class CalibrationManagementTests(unittest.TestCase):
    def test_normaliza_codigo_ignorando_hifen_espaco_e_caixa(self):
        self.assertEqual(normalizar_codigo("E 10060"), "E10060")
        self.assertEqual(normalizar_codigo("e-10060"), "E10060")
        self.assertEqual(normalizar_codigo("PAQ-001"), "PAQ001")

    def test_classifica_vencimento_nos_limites(self):
        hoje = date(2026, 10, 16)
        self.assertEqual(classificar_vencimento(date(2026, 10, 15), hoje), "VENCIDO")
        self.assertEqual(classificar_vencimento(date(2026, 10, 16), hoje), "A VENCER")
        self.assertEqual(classificar_vencimento(date(2026, 11, 15), hoje), "A VENCER")
        self.assertEqual(classificar_vencimento(date(2026, 11, 16), hoje), "NO PRAZO")
        self.assertEqual(
            classificar_vencimento(None, hoje),
            "SEM PRÓXIMA CALIBRAÇÃO",
        )

    def test_validates_required_headers_using_names(self):
        self.assertTrue(
            validar_estrutura_cal(
                [
                    "NÚMERO\nEQUIPAMENTO",
                    "EQUIPAMENTO",
                    "PRÓXIMA CALIBRAÇÃO",
                    "STATUS CALIBRAÇÃO",
                ]
            )
        )
        self.assertFalse(validar_estrutura_cal(["Nº EQUIPAMENTO", "DESCRIÇÃO"]))

    def test_loader_ignores_formulas_and_recalculates_dates_and_status(self):
        dataframe = carregar_calibracao(criar_planilha_calibracao())
        self.assertEqual(len(dataframe), 1)
        self.assertEqual(dataframe.loc[0, "Nº Equipamento"], "E-10060")
        self.assertEqual(dataframe.loc[0, "Próxima Calibração"].date(), date(2026, 1, 1))
        self.assertEqual(dataframe.loc[0, "Status Calibração"], "VENCIDO")
        self.assertEqual(dataframe.loc[0, "Fabricante"], "Não Definido")
        self.assertEqual(dataframe.loc[0, "Código Normalizado"], "E10060")

    def test_loader_errors_when_required_structure_is_missing(self):
        workbook = openpyxl.Workbook()
        worksheet = workbook.active
        worksheet.title = "CAL"
        worksheet.append(["Nº EQUIPAMENTO", "DESCRIÇÃO", "STATUS CALIBRAÇÃO"])
        worksheet.append(["A1", "Balança", "NO PRAZO"])
        buffer = BytesIO()
        workbook.save(buffer)
        with self.assertRaises(ErroEstruturaCalibracao):
            carregar_calibracao(buffer.getvalue())

    def test_exports_to_excel_and_pdf(self):
        dataframe = pd.DataFrame(
            [
                {
                    "Nº Equipamento": "E10060",
                    "Descrição": "Paquímetro",
                    "Unidade": "mm",
                    "Capacidade": "150",
                    "Local de Uso": "Laboratório",
                    "Próxima Calibração": pd.Timestamp("2026-10-01"),
                    "Status Calibração": "VENCIDO",
                    "Fabricante": "Não Definido",
                    "Setor": "Qualidade",
                    "Status Equipamento": "ATIVO",
                }
            ]
        )
        excel = exportar_excel_calibracao(dataframe)
        loaded = openpyxl.load_workbook(BytesIO(excel), data_only=True)
        self.assertEqual(loaded.active["A2"].value, "E10060")
        loaded.close()
        pdf = PdfReader(BytesIO(_pdf_calibracao(dataframe)))
        self.assertEqual(len(pdf.pages), 1)
        self.assertIn("E10060", pdf.pages[0].extract_text())

    def test_categorical_charts_use_distinct_qmax_colors(self):
        dataframe = pd.DataFrame(
            {
                "Setor": ["Qualidade", "Manutenção", "Produção"],
                "Quantidade": [10, 8, 5],
            }
        )
        figure = _figura_barras_categoricas(dataframe, "Setor")
        colors = [trace.marker.color for trace in figure.data]
        self.assertEqual(len(colors), 3)
        self.assertEqual(len(set(colors)), 3)
        self.assertTrue(set(colors).issubset(PALETA_GRAFICOS_CALIBRACAO))
        self.assertFalse(figure.layout.showlegend)

    def test_monthly_due_chart_is_a_status_colored_time_series(self):
        dataframe = pd.DataFrame(
            [
                {
                    "Próxima Calibração": pd.Timestamp("2026-10-01"),
                    "Status Calibração": "VENCIDO",
                },
                {
                    "Próxima Calibração": pd.Timestamp("2026-10-22"),
                    "Status Calibração": "A VENCER",
                },
                {
                    "Próxima Calibração": pd.Timestamp("2026-12-15"),
                    "Status Calibração": "NO PRAZO",
                },
            ]
        )
        figure = _figura_vencimentos_mensais(dataframe)
        self.assertTrue(all("lines" in trace.mode and "markers" in trace.mode for trace in figure.data))
        self.assertEqual(
            {trace.name: trace.line.color for trace in figure.data},
            {
                "VENCIDO": CORES_STATUS_CALIBRACAO["VENCIDO"],
                "A VENCER": CORES_STATUS_CALIBRACAO["A VENCER"],
                "NO PRAZO": CORES_STATUS_CALIBRACAO["NO PRAZO"],
            },
        )

    def test_report_status_filters_match_calculated_status_values(self):
        dataframe = pd.DataFrame(
            {
                "Nº Equipamento": ["E1", "E2", "E3"],
                "Status Calibração": ["VENCIDO", "A VENCER", "NO PRAZO"],
            }
        )
        self.assertEqual(
            _filtrar_status_relatorio(dataframe, "Vencidos")["Nº Equipamento"].tolist(),
            ["E1"],
        )
        self.assertEqual(
            _filtrar_status_relatorio(dataframe, "A Vencer")["Nº Equipamento"].tolist(),
            ["E2"],
        )
        self.assertEqual(
            _filtrar_status_relatorio(dataframe, "No Prazo")["Nº Equipamento"].tolist(),
            ["E3"],
        )


if __name__ == "__main__":
    unittest.main()
